"""No-overwrite XLSX snapshots with numerical readback verification."""
from __future__ import annotations

from datetime import date, datetime
import hashlib
import json
import math
import os
from pathlib import Path
import tempfile

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from shared_code.reader import DataContractError
from .assessments import BLOCK_FEATURES, DAILY_FEATURES, FITROCKR_FEATURES, METHOD_ID, UNITS


def cell_value(value, column):
    if value is None or (not isinstance(value, (list, dict)) and pd.isna(value)):
        return None
    if isinstance(value, np.generic): value = value.item()
    if isinstance(value, float) and not math.isfinite(value): return None
    if isinstance(value, pd.Timestamp): return value.to_pydatetime()
    if isinstance(value, (dict, list)): return json.dumps(value, sort_keys=True, default=str)
    if isinstance(value, str):
        if len(value) > 32767:
            raise DataContractError("xlsx_cell_too_long", f"Cell exceeds Excel limit: {column}")
        if value == "": return None
        if column in {"date", "anchor_date", "T0_anchor", "T1_anchor"}: return date.fromisoformat(value)
    return value


def add_table(wb, name, frame):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = True
    ws.freeze_panes = "D2" if len(frame.columns) > 3 else "B2"
    for i, col in enumerate(frame.columns, 1):
        c = ws.cell(1, i, str(col))
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 18 if col in {"timepoint", "anchor_date", "date", "cycle"} else 24
    ws.row_dimensions[1].height = 32
    for nr, record in enumerate(frame.itertuples(index=False, name=None), 2):
        for nc, (column, value) in enumerate(zip(frame.columns, record), 1):
            value = cell_value(value, column)
            c = ws.cell(nr, nc, value)
            if isinstance(value, str): c.data_type = "s"
            if isinstance(value, date): c.number_format = "yyyy-mm-dd hh:mm:ss.000" if isinstance(value, datetime) else "yyyy-mm-dd"
        ws.row_dimensions[nr].height = 15
    if len(frame):
        table = Table(displayName=name.replace("_", ""), ref=f"A1:{get_column_letter(len(frame.columns))}{len(frame)+1}")
        table.tableStyleInfo = TableStyleInfo(name=None, showRowStripes=False, showColumnStripes=False)
        ws.add_table(table)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 0, 0
    ws.print_title_rows = "1:1"
    if len(frame.columns) > 3: ws.print_title_cols = "A:C"
    return ws


def export_assessment(destination, tables, plans, reports, clinical_issues, method, *, writer=None):
    destination = Path(destination)
    if destination.exists(): raise DataContractError("output_exists", f"Workbook exists: {destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    # Selected feature names are the only physiological outputs in the detailed
    # view. Their calculation evidence and source IDs remain separate.
    blocks = tables["Blocks_20_per_T"]
    allowed = {"participant_key", "timepoint", "anchor_date", "date_status", "relative_day", "date", "cycle", "circadian_cycle", "window_start", "window_end",
               "n_beats", "n_adjacent_pairs", "health_coverage_pct", "wearing_time_pct_observed", "worn_pct", "hrv_segments", "hrv_longest_segment_s", "hrv_pairs_excluded", "hrv_clock_mismatch_pairs"}
    sources = {*BLOCK_FEATURES.values(), *DAILY_FEATURES.values(), *filter(None, FITROCKR_FEATURES.values())}
    allowed |= sources
    allowed |= {f"{c}{suffix}" for c in sources for suffix in ("_review", "_n", "_n_epochs", "_status")}
    tables = {**tables, "Blocks_20_per_T": blocks[[c for c in blocks if c in allowed]]}
    anchors = [{k: p[k] for k in ("participant_key", "timepoint", "anchor_date", "date_status", "anchor_basis")} | {
        "source_cells_json": json.dumps(p["source_cells"]), "override_json": json.dumps(p["override"]) if p["override"] else None,
        "window_basis": p["window_basis"], "window_json": json.dumps(p["window"]) if p.get("window") else None, "dates_json": json.dumps(p["dates"])} for p in plans]
    qc = []
    for report in reports:
        for name, stream in report["streams"].items():
            for finding in stream["findings"]:
                qc.append({"participant_key": report["participant_key"], "timepoint": report["timepoint"], "variant": report["variant"],
                           "stream": name, **{k: finding[k] for k in ("code", "feature", "severity", "count", "rule")}, "examples_json": json.dumps(finding["examples"])})
    for finding in clinical_issues:
        qc.append({"participant_key": plans[0]["participant_key"], "timepoint": finding.get("timepoint"), "variant": "clinical", "stream": "clinical",
                   "code": finding["code"], "severity": "review", "count": 1, "rule": finding.get("details"), "examples_json": json.dumps(finding, default=str)})
    tables["Anchors"] = pd.DataFrame(anchors)
    tables["QC"] = pd.DataFrame(qc, columns=["participant_key", "timepoint", "variant", "stream", "code", "feature", "severity", "count", "rule", "examples_json"])
    definitions = [
        ("method_id", METHOD_ID, None),
        ("BBI", "Raw BBI only (Enhanced BBI fills detection gaps and is not read); inclusive 230-1550 ms screen; wear>0 and not MET>=3 (unknown MET kept); beats >20% from the 31-beat median excluded; pairs only when the UTC step equals the following BBI within 50 ms", None),
        ("second_beat_stream", "5-min bins where raw BBI sum exceeds 1.05 x clock time (a second wearer), plus one bin either side: BBI, stress and minute HR excluded", None),
        ("zero_coded_missing", "Fitrockr writes 0 where it computed nothing. Excluded: MET 0; MET quarters with missing calories; worn 0-kcal records (total, active and steps together); total kcal on days without resting energy; steps when none were ever recorded", None),
        ("Summary_Fitrockr_HRV", "Same as Summary except RMSSD, SDNN, LF/HF and SD2/SD1 = block means of Fitrockr's 5-min values (SD2/SD1 = 1/SD1_SD2); SampEn not provided by Fitrockr", None),
        ("window", "Xsens assessment date T +/-2 local calendar days; night 00-06, morning 06-12, afternoon 12-18, evening 18-24; half-open blocks", None),
        ("averaging", "Compute each daily six-hour feature, then average finite daily values with equal day weight; no duration/observation-count weighting across days", None),
        ("contributing_days", "Feature-specific count of valid days; fewer than 3 means the T value is blank; unresolved anchors have blank counts", "days"),
        ("block_wear", "A six-hour block-day counts only when worn for at least 50% of it (3 h); otherwise all its block features are blank", None),
        ("T1_window", "The watch is handed over at T0 and paired days later. T1 starts on the later of Xsens day -2 and the first wholly worn day (the day after the first worn health interval, or that day if it starts at 00:00), keeps up to 5 days and ends before T2's Xsens day (Anchors/window_basis t1_first_worn_day)", None),
        ("window_override", "An explicit start-end window (Anchors/window_json) replaces Xsens date +/-2 and the T1 rule for that T", None),
        ("advisory_review_days", "Count of contributing daily values carrying existing review warnings; these values remain in this exploratory output", "days"),
        ("Daily", "RHR and sleep counted once per date; daily values repeat on the 20-row view for context, not independent observations", None),
        ("calories", "Total Calories (resting + active), mean per eligible source interval. Existing high-rate warnings remain; upper-limit decision deferred", "kcal per source interval"),
        ("zero_HR", "RHR requires positive finite values. Physiological stress requires a positive finite HR reference. No separate mean-HR feature", None),
        ("HRV", "SDNN sample standard deviation; RMSSD and SD2/SD1 use supported original adjacent pairs; no concatenation across removed observations", None),
        ("advanced_HRV", "LF/HF: median of finite complete 300s epoch estimates; SampEn: median of finite 1000-contiguous-beat epochs. Existing math/parameters retained", None),
        ("gap_marker", "Blank gap duration is expected except on gap records; positive or invalid populated gap duration breaks continuity", None),
        ("T0_T1", "Exploratory only; reports shared days because five-day windows may overlap. T0 is excluded from the primary T1-T10 Summary", None),
        ("QC_counts", "Counts belong to the named assessment/source run. Windows and rules can overlap; do not sum as unique affected source observations", None),
        ("provenance", "Companion metadata.json contains source hashes, clinical cells, per-run counts and workbook verification; method.json contains configuration and implementation hashes", None),
        ("validation_limit", "Implementation verification is not validation of sinus rhythm, stationarity, ECG agreement, clinical interpretation or rehabilitation efficacy", None),
        ("source_timing_reference", "https://www.fitrockr.com/help-center/data-dictionary/", None),
        ("source_preference_reference", "https://www.fitrockr.com/help-center/hrv/", None),
    ]
    for feature, source in {**BLOCK_FEATURES, **DAILY_FEATURES}.items():
        definitions.append((feature, f"Derived from {source}; available-day means in Summary; support and reviews in Support", UNITS[feature]))
    tables["Definitions"] = pd.DataFrame(definitions, columns=["field", "definition", "unit"])
    order = ["Summary", *(["Summary_Fitrockr_HRV"] if "Summary_Fitrockr_HRV" in tables else []), "Blocks_20_per_T", "Daily", "Support", "T0_T1", "Anchors", "QC", "HRV_epochs", "Definitions"]
    fd, temporary = tempfile.mkstemp(prefix=".garmin-", suffix=".xlsx", dir=destination.parent)
    os.close(fd)
    try:
        if writer is None:
            wb = Workbook()
            wb.remove(wb.active)
            for name in order: add_table(wb, name, tables[name])
            wb.save(temporary)
        else:
            # Optional authoring adapter; numerical and shape verification below
            # is independent of the writer. The default stays portable Python.
            writer(Path(temporary), {name: tables[name] for name in order})
        readback = load_workbook(temporary, read_only=True, data_only=True)
        checked = 0
        for name, frame in tables.items():
            ws = readback[name]
            if ws.max_row is None or ws.max_column is None:
                ws.calculate_dimension(force=True)
            if ws.max_row != len(frame)+1 or ws.max_column != len(frame.columns):
                raise DataContractError("xlsx_shape_mismatch", f"Readback shape differs: {name}: {(ws.max_row, ws.max_column)} versus {(len(frame)+1, len(frame.columns))}")
            for expected, actual in zip(frame.itertuples(index=False, name=None), ws.iter_rows(min_row=2, values_only=True), strict=True):
                for col, value, observed in zip(frame.columns, expected, actual, strict=True):
                    value = cell_value(value, col)
                    if isinstance(value, float): equal = observed is not None and math.isclose(value, observed, rel_tol=1e-13, abs_tol=1e-10)
                    elif isinstance(value, date) and not isinstance(value, datetime): equal = observed is not None and observed.date() == value
                    elif isinstance(value, datetime):
                        # Excel stores dates to millisecond precision.
                        equal = observed is not None and abs((value-observed).total_seconds()) <= .0005
                    else: equal = value == observed
                    if not equal: raise DataContractError("xlsx_value_mismatch", f"Readback value differs: {name}/{col}")
                    checked += 1
        readback.close()
        # Atomic no-replace publication; a concurrent creator cannot be overwritten.
        os.link(temporary, destination)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return {"sha256": hashlib.sha256(destination.read_bytes()).hexdigest(), "cells_checked": checked,
            "sheet_rows": {name: len(frame) for name, frame in tables.items()}, "numerical_readback_passed": True}
