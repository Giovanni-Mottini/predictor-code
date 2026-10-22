from __future__ import annotations

from datetime import datetime
import hashlib
import io
import json
from pathlib import Path
import warnings
from zipfile import ZipFile, ZipInfo

import openpyxl
import pytest

from predictor_import.adapter import make_plan, write_plan
from predictor_import.inspection import XSENS_SHEETS
from predictor_import.inspection import normalize_hemisphere
from predictor_import.inspection import inspect_file
from predictor_import.inventory import scan
from predictor_import.naming import filename_subject, parse_xsens_name
from predictor_import.inventory_output import write_inventory
from predictor_import.__main__ import main


def workbook(path: Path, *, metadata_id: str | None = None, hemisphere="Right") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    book.remove(book.active)
    if metadata_id:
        sheet = book.create_sheet("TIMELINE_01")
        sheet.append(["Hemisphere", "Dominant Hand"])
        sheet.append([hemisphere, "left"])
        sheet.append([metadata_id])
        sheet.append(["T1_d1", "T1_d2", "T5"])
        sheet.append([datetime(2030, 1, 1), datetime(2030, 1, 2), datetime(2030, 2, 1)])
    else:
        for title in XSENS_SHEETS:
            sheet = book.create_sheet(title)
            sheet.append(["Frame", "value"])
            sheet.append([1, 0.0])
    book.save(path)
    return path


@pytest.fixture
def source(tmp_path):
    root = tmp_path / "Dati PREDICTOR"
    for cohort in ("SANI", "PAZIENTI", "CONTROLLO PAZIENTI"):
        (root / cohort).mkdir(parents=True)
    return root


@pytest.mark.parametrize("name,expected", [
    ("FG91_00_TASK1_T0.xlsx", ("T0", "00", "task1", None)),
    ("fg91_1_task2_t10_d3.XLSX", ("T10", "1", "task2", None)),
    ("BA0101XX00-t1_task1_dx.xlsx", ("T1", "dx", "task1", None)),
    ("BA0202YY00_t1_mot2_sx.xlsx", ("T1", "sx", "mot2", None)),
    ("BA0303ZZ00_task4_t1_dx.xlsx", ("T1", "dx", "task4", None)),
    ("BA0404WW00_task3_rip5_sx.xlsx", (None, "sx", "task3", 5)),
    ("BA0404WW00_task1_dx.xlsx", (None, "dx", "task1", None)),
    ("ctrlfg91_01_b&b_t1.xlsx", ("T1", "01", "box_and_blocks", None)),
])
def test_explicit_filename_variants(name, expected):
    parsed = parse_xsens_name(name)
    assert (parsed["timepoint"], parsed["original_limb_code"], parsed["activity"], parsed["repetition"]) == expected


@pytest.mark.parametrize("name", ["FG91_VAL.xlsx", "FG91_00_TASK5_T1.xlsx", "FG91_00_TASK1_T1.xlsx.partial", "random.xlsx"])
def test_unknown_names_are_not_guessed(name):
    assert parse_xsens_name(name) is None


def test_folder_aliases_are_cohort_specific():
    assert filename_subject("SANO_FG91", "SANI") == "FG91"
    assert filename_subject("SANO_FG91", "PAZIENTI") == "SANO_FG91"
    assert filename_subject("CTRLFG_91", "CONTROLLO PAZIENTI") == "CTRLFG91"
    assert filename_subject("FG95_dropout", "PAZIENTI") == "FG95"


def test_multi_gigabyte_csv_member_still_gets_bounded_header_inspection(tmp_path, monkeypatch):
    info = ZipInfo("bbi.csv")
    info.file_size, info.compress_size, info.CRC = 3 * 2**30, 123, 0
    reads = []

    class BoundedStream(io.BytesIO):
        def read(self, size=-1):
            reads.append(size)
            assert 0 < size <= 65536
            return super().read(size)

    class Archive:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def infolist(self):
            return [info]

        def open(self, member):
            return BoundedStream(b"Time(ms),BBI(ms)\n1,800\n")

    monkeypatch.setattr("predictor_import.inspection.ZipFile", lambda _: Archive())
    result = inspect_file(tmp_path / "large.zip")
    member = result["archive_members"][0]
    assert member["size_bytes"] == 3 * 2**30
    assert member["inspection_status"] == "header_readable"
    assert member["columns"] == ["Time(ms)", "BBI(ms)"]
    assert reads == [65536]


@pytest.mark.parametrize("text,expected", [
    ("left frontal", "left"), ("right occipital", "right"),
    ("left cerebellar", "left"), ("Fronto-parietale dx", "right"),
    ("left fronto-temporal", "left"),
    ("left and right", None), ("not right", None), ("possible left", None),
])
def test_hemisphere_description_has_explicit_unambiguous_side(text, expected):
    assert normalize_hemisphere(text) == expected


def test_patient_metadata_and_exact_dates(source):
    workbook(source / "PAZIENTI/FG91/FG91.xlsx", metadata_id="FG91")
    workbook(source / "PAZIENTI/FG91/Xsens/T1/FG91_01_TASK1_T1.xlsx")
    workbook(source / "PAZIENTI/FG91/Xsens/T5/FG91_00_TASK2_T5.xlsx")
    plan = make_plan(scan(source), "PAZIENTI/FG91")
    metadata = plan["subjects"][0]["subject"]["clinical_metadata"]
    assert metadata["paretic_side"] == "left"
    assert metadata["non_paretic_side"] == "right"
    assert metadata["dominant_hand"] == "left"
    records = {r["timepoint"]: r for r in plan["subjects"][0]["manifest"]["records"] if r["timepoint"]}
    assert records["T1"]["anatomical_side"] == "left"
    assert records["T1"]["acquisition_date"] is None
    assert records["T5"]["acquisition_date"] == "2030-02-01"
    assert records["T5"]["limb_condition"] == "non_paretic"
    assert all(not r["ready_for_features"] for r in records.values())


def test_explicit_day_uses_exact_timeline_date(source):
    workbook(source / "PAZIENTI/FG91/FG91_VAL.xlsx", metadata_id="FG91")
    workbook(source / "PAZIENTI/FG91/Xsens/FG91_1_TASK1_T1_d2.xlsx")
    record = next(r for r in make_plan(scan(source))["subjects"][0]["manifest"]["records"] if r["activity"])
    assert record["session_id"] == "T1_d2"
    assert record["acquisition_date"] == "2030-01-02"


def test_explicit_protocol_days_are_never_selected_as_duplicate_exports(source):
    workbook(source / "PAZIENTI/FG91/Xsens/FG91_1_TASK1_T1_d1.xlsx")
    workbook(source / "PAZIENTI/FG91/Xsens/FG91_1_TASK1_T1_d2.xlsx")
    records = make_plan(scan(source))["subjects"][0]["manifest"]["records"]
    assert len(records) == 2
    assert all(r["disposition"] == "mapped" for r in records)
    assert {r["session_id"] for r in records} == {"T1_d1", "T1_d2"}
    assert all(r["selection_assumption"] is None for r in records)


def test_wrong_subject_inside_evaluation_is_not_used(source):
    workbook(source / "PAZIENTI/FG91/FG91_VAL.xlsx", metadata_id="FG93")
    workbook(source / "PAZIENTI/FG91/Xsens/FG91_01_TASK1_T1.xlsx")
    adapted = make_plan(scan(source))["subjects"][0]
    assert adapted["subject"]["clinical_metadata"]["status"] == "needs_review"
    assert all(r["anatomical_side"] is None for r in adapted["manifest"]["records"])


def test_conflicting_evaluations_are_not_merged(source):
    workbook(source / "PAZIENTI/FG91/FG91_VAL.xlsx", metadata_id="FG91", hemisphere="Left")
    workbook(source / "PAZIENTI/FG91/FG91.xlsx", metadata_id="FG91", hemisphere="Right")
    subject = make_plan(scan(source))["subjects"][0]["subject"]
    assert subject["clinical_metadata"]["status"] == "needs_review"
    assert subject["clinical_metadata"]["paretic_side"] is None


def test_sani_missing_evaluation_is_expected_and_identity_is_distinct(source):
    workbook(source / "SANI/SANO_FG91/XSens/FG91_t1_task1_dx.xlsx")
    workbook(source / "PAZIENTI/FG91/Xsens/FG91_00_TASK1_T1.xlsx")
    plan = make_plan(scan(source))
    subjects = {p["subject"]["participant_key"]: p for p in plan["subjects"]}
    assert set(subjects) == {"SANI/SANO_FG91", "PAZIENTI/FG91"}
    sani = subjects["SANI/SANO_FG91"]
    assert sani["subject"]["clinical_metadata"]["status"] == "not_collected_by_protocol"
    record = sani["manifest"]["records"][0]
    assert record["anatomical_side"] == "right"
    assert record["limb_condition"] == "non_paretic"
    assert record["timepoint_kind"] == "session_label"


def test_control_codes_follow_confirmed_clinical_interpretation(source):
    workbook(source / "CONTROLLO PAZIENTI/CTRLFG_91/XSENS/ctrlfg91_01_task1_t1.xlsx")
    record = make_plan(scan(source))["subjects"][0]["manifest"]["records"][0]
    assert record["limb_condition"] == "paretic"
    assert record["anatomical_side"] is None
    assert record["disposition"] == "mapped"
    assert record["normalized_target"] is not None


def test_identity_modality_and_timepoint_conflicts(source):
    workbook(source / "PAZIENTI/FG94/MOTORE/T5/fg93_01_task3_t6.xlsx")
    inventory = scan(source)
    file = inventory["files"][0]
    assert {"subject_identity_conflict", "modality_folder_conflict", "timepoint_folder_conflict"}.issubset(file["issues"])
    record = make_plan(inventory)["subjects"][0]["manifest"]["records"][0]
    assert record["disposition"] == "excluded"
    assert record["anatomical_side"] is None


def test_full_and_repetition_exports_cannot_be_double_counted(source):
    workbook(source / "SANI/SANO_BA001/XSens/Export/BA001_task1_dx.xlsx")
    workbook(source / "SANI/SANO_BA001/XSens/Export/BA001_task1_rip1_dx.xlsx")
    inventory = scan(source)
    assert all("overlapping_recording_representations" in f["issues"] for f in inventory["files"])
    selected = [f for f in inventory["files"] if f["disposition"] == "mapped"]
    ignored = [f for f in inventory["files"] if f["disposition"] == "excluded"]
    assert len(selected) == len(ignored) == 1
    assert selected[0]["parsed_name"]["representation"] == "full_recording"
    assert ignored[0]["duplicate_of_source_id"] == selected[0]["source_id"]
    assert selected[0]["selection_assumption"] == "same_acquisition_pending_professor_verification"


def test_distinct_individual_repetitions_are_not_duplicates(source):
    for rep in (1, 2):
        workbook(source / f"SANI/SANO_BA001/XSens/Export/BA001_task1_rip{rep}_dx.xlsx")
    assert all("multiple_recordings_same_key" not in f["issues"] for f in scan(source)["files"])


def test_zip_headers_are_read_without_extraction_and_overlap_is_flagged(source):
    garmin = source / "PAZIENTI/FG91/Garmin"
    garmin.mkdir(parents=True)
    with ZipFile(garmin / "fg91.zip", "w") as archive:
        archive.writestr("fg91/daily-health-log.csv", "Calendar Date (UTC),Total Calories (kcal)\n2030-01-01,42\n")
    (garmin / "daily-health-log.csv").write_text("Calendar Date (UTC),Total Calories (kcal)\n2030-01-01,42\n")
    inventory = scan(source)
    archive = next(f for f in inventory["files"] if f["artifact_kind"] == "garmin_archive")
    assert archive["archive_members"][0]["columns"] == ["Calendar Date (UTC)", "Total Calories (kcal)"]
    assert all("archive_and_loose_table_overlap_candidate" in f["issues"] for f in inventory["files"])
    assert not (garmin / "fg91").exists()


@pytest.mark.parametrize("member", ["../outside.csv", "/absolute.csv", "C:/outside.csv", "folder\\outside.csv"])
def test_unsafe_zip_paths_are_reported(source, member):
    path = source / "SANI/SANO_BA001/Garmin/export.zip"
    path.parent.mkdir(parents=True)
    with ZipFile(path, "w") as archive:
        archive.writestr(member, "value\n1\n")
    file = scan(source)["files"][0]
    assert file["disposition"] == "excluded"
    assert not file["archive_members"][0]["safe_path"]


def test_duplicate_archive_member_names_are_blocked(source):
    path = source / "SANI/SANO_BA001/Garmin/export.zip"
    path.parent.mkdir(parents=True)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        with ZipFile(path, "w") as archive:
            archive.writestr("table.csv", "value\n1\n")
            archive.writestr("table.csv", "value\n2\n")
    file = scan(source)["files"][0]
    assert all(not member["unique_path"] for member in file["archive_members"])
    assert "unsafe_or_duplicate_archive_member" in file["issues"]


def test_partial_and_corrupt_files_are_distinguished(source):
    folder = source / "PAZIENTI/FG91/Xsens"
    folder.mkdir(parents=True)
    (folder / "FG91_01_TASK1_T1.xlsx.abc.partial").write_bytes(b"partial download")
    (folder / "broken.xlsx").write_bytes(b"not an XLSX package")
    files = {Path(f["source_path"]).name: f for f in scan(source)["files"]}
    assert files["FG91_01_TASK1_T1.xlsx.abc.partial"]["disposition"] == "deferred"
    assert files["broken.xlsx"]["inspection_status"] == "inspection_error"
    assert files["broken.xlsx"]["disposition"] == "excluded"


def test_no_source_files_are_modified_and_reports_are_complete(source, tmp_path):
    workbook(source / "PAZIENTI/FG91/Xsens/FG91_01_TASK1_T1.xlsx")
    (source / "SANI/SANO_EMPTY").mkdir()
    def fingerprints():
        return {p.relative_to(source).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in source.rglob("*") if p.is_file()}
    before = fingerprints()
    inventory = scan(source)
    report = tmp_path / "report"
    write_inventory(inventory, report)
    write_plan(make_plan(inventory), tmp_path / "standardized", hash_sources=True)
    assert fingerprints() == before
    assert inventory["summary"]["empty_subject_count"] == 1
    assert {p.name for p in report.iterdir()} == {"inventory.json", "summary.json", "files.csv", "subjects.csv", "issues.csv"}


def test_source_change_since_inventory_is_deferred(source, tmp_path):
    path = workbook(source / "PAZIENTI/FG91/Xsens/FG91_01_TASK1_T1.xlsx")
    plan = make_plan(scan(source))
    path.write_bytes(path.read_bytes() + b"changed")
    report = write_plan(plan, tmp_path / "standardized")
    assert report["deferred_sources"]
    record = json.loads((tmp_path / "standardized/PAZIENTI/FG91/manifest.json").read_text())["records"][0]
    assert record["disposition"] == "deferred"


def test_changed_evaluation_clears_derived_sides(source, tmp_path):
    path = workbook(source / "PAZIENTI/FG91/FG91_VAL.xlsx", metadata_id="FG91")
    workbook(source / "PAZIENTI/FG91/Xsens/FG91_01_TASK1_T1.xlsx")
    plan = make_plan(scan(source))
    path.write_bytes(path.read_bytes() + b"changed")
    output = tmp_path / "standardized"
    write_plan(plan, output)
    subject = json.loads((output / "PAZIENTI/FG91/subject.json").read_text())
    record = next(r for r in json.loads((output / "PAZIENTI/FG91/manifest.json").read_text())["records"] if r["activity"])
    assert subject["clinical_metadata"]["status"] == "needs_review"
    assert record["anatomical_side"] is None


def test_staging_is_repeatable_and_adds_optional_content_hash(source, tmp_path):
    path = workbook(source / "PAZIENTI/FG91/Xsens/FG91_01_TASK1_T1.xlsx")
    inventory = scan(source)
    output = tmp_path / "standardized"
    write_plan(make_plan(inventory), output, hash_sources=True)
    first = (output / "PAZIENTI/FG91/manifest.json").read_bytes()
    write_plan(make_plan(inventory), output, hash_sources=True)
    assert first == (output / "PAZIENTI/FG91/manifest.json").read_bytes()
    assert json.loads(first)["records"][0]["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()


def test_output_cannot_be_written_into_raw_sources(source):
    workbook(source / "PAZIENTI/FG91/Xsens/FG91_01_TASK1_T1.xlsx")
    inventory = scan(source)
    with pytest.raises(ValueError):
        write_plan(make_plan(inventory), source / "standardized")
    with pytest.raises(ValueError):
        write_inventory(inventory, source / "reports")


def test_source_symlinks_are_not_followed(source, tmp_path):
    target = tmp_path / "outside.csv"
    target.write_text("value\n123\n")
    folder = source / "SANI/SANO_BA001/Garmin"
    folder.mkdir(parents=True)
    (folder / "linked.csv").symlink_to(target)
    inventory = scan(source)
    assert inventory["files"][0]["artifact_kind"] == "symlink"
    assert inventory["files"][0]["disposition"] == "excluded"


def test_unknown_participant_filter_fails(source):
    with pytest.raises(ValueError):
        make_plan(scan(source), "PAZIENTI/DOES_NOT_EXIST")


def test_unknown_schema_version_is_rejected(source):
    inventory = scan(source)
    inventory["schema_version"] = "9.9.9"
    with pytest.raises(ValueError, match="schema version"):
        make_plan(inventory)


def test_cli_dry_run_produces_plan_only(source, tmp_path):
    workbook(source / "SANI/SANO_BA001/XSens/BA001_t1_task1_dx.xlsx")
    output = tmp_path / "report"
    assert main(["inventory", str(source), "--output", str(output)]) == 0
    assert main(["adapt", str(output / "inventory.json"), "--output", str(output)]) == 0
    plan = json.loads((output / "adapter_plan.json").read_text())
    assert plan["summary"]["subjects"] == 1
    assert not (output / "SANI").exists()


def test_contract_field_sets_match_machine_schemas(source):
    workbook(source / "PAZIENTI/FG91/FG91_VAL.xlsx", metadata_id="FG91")
    workbook(source / "PAZIENTI/FG91/Xsens/FG91_01_TASK1_T1.xlsx")
    inventory = scan(source)
    adapted = make_plan(inventory)["subjects"][0]
    schemas = Path(__file__).resolve().parents[2] / "schemas"
    subject_schema = json.loads((schemas / "subject.schema.json").read_text())
    manifest_schema = json.loads((schemas / "manifest.schema.json").read_text())
    inventory_schema = json.loads((schemas / "inventory.schema.json").read_text())
    assert set(adapted["subject"]) == set(subject_schema["required"])
    assert set(adapted["manifest"]) == set(manifest_schema["required"])
    assert set(inventory) == set(inventory_schema["required"])
    for record in adapted["manifest"]["records"]:
        assert set(record) == set(manifest_schema["$defs"]["record"]["required"])
    from jsonschema import Draft202012Validator
    for schema, document in [(subject_schema, adapted["subject"]), (manifest_schema, adapted["manifest"]), (inventory_schema, inventory)]:
        Draft202012Validator.check_schema(schema)
        Draft202012Validator(schema).validate(document)


def test_sani_simulation_is_distinct_from_actual_paresis(source):
    workbook(source / "SANI/SANO_BI001/XSens/BI001_task1_t1_sim.xlsx")
    record = make_plan(scan(source))["subjects"][0]["manifest"]["records"][0]
    assert record["limb_condition"] == "simulated"
    assert record["anatomical_side"] is None
    assert record["simulation_types_by_repetition"] is None
    assert record["expected_repetitions"] == 4
    assert record["expected_marker_count"] == 5
    assert record["disposition"] == "mapped"


def test_motore_and_xsens_mot_activities_are_ignored(source):
    workbook(source / "SANI/SANO_BA001/XSens/BA001_t1_mot1_dx.xlsx")
    folder = source / "SANI/SANO_BA001/MOTORE"
    folder.mkdir()
    (folder / "stats.json").write_text('{}')
    inventory = scan(source)
    assert all(f["disposition"] == "excluded" for f in inventory["files"])
    assert all(f["selection_reason"] == "motore_out_of_scope" for f in inventory["files"])
    assert len(make_plan(inventory)["subjects"][0]["manifest"]["records"]) == 2
    assert inventory["summary"]["data_file_count"] == 2


def test_archive_preference_does_not_hide_loose_tables_missing_from_archive(source):
    folder = source / "PAZIENTI/FG91/Garmin"
    folder.mkdir(parents=True)
    with ZipFile(folder / "fg91.zip", "w") as archive:
        archive.writestr('table.csv', 'value\n1\n')
    (folder / "table.csv").write_text('value\n1\n')
    (folder / "extra.csv").write_text('value\n2\n')
    files = {Path(f['source_path']).name:f for f in scan(source)['files']}
    assert files['fg91.zip']['disposition'] == 'mapped'
    assert files['fg91.zip']['selected_archive_members'] == ['table.csv']
    assert files['table.csv']['disposition'] == 'excluded'
    assert files['extra.csv']['disposition'] == 'mapped'
