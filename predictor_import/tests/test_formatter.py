from __future__ import annotations

from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile

from jsonschema import Draft202012Validator
import openpyxl
import pyarrow.parquet as pq
import pytest

from predictor_import.adapter import make_plan, write_plan
from predictor_import.formatter import format_subject
from predictor_import.inspection import XSENS_SHEETS
from predictor_import.inventory import scan


def acquisition(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name in XSENS_SHEETS:
        sheet = book.create_sheet(name)
        if name == "General Information":
            sheet.append(["Frame Rate", 60])
            sheet.append(["Recorded Date UTC", datetime(2030, 1, 2, 12, 30)])
            sheet.append(["Original File Name", "original.mvn"])
        elif name == "Markers":
            sheet.append(["Frame", "Marker Text"])
            for frame in range(1, 6):
                sheet.append([frame, f"Marker_{frame}"])
        else:
            sheet.append(["Frame", "Right Hand x", "Left Hand x"])
            for frame in range(7):
                sheet.append([frame, -0.1234567890123456 * frame, None if frame == 3 else frame * 0.1])
    book.save(path)


@pytest.fixture
def staged(tmp_path):
    root = tmp_path / "raw"
    source = root / "SANI/SANO_BA001/XSens/BA001_t1_task1_dx.xlsx"
    acquisition(source)
    output = tmp_path / "standardized"
    inventory = scan(root)
    write_plan(make_plan(inventory), output)
    return source, output / "SANI/SANO_BA001", inventory


def descriptor(subject):
    manifest = json.loads((subject / "manifest.json").read_text())
    record = manifest["records"][0]
    directory = subject / record["normalized_target"]
    return record, directory, json.loads((directory / "recording.json").read_text())


def test_every_sheet_and_value_is_preserved(staged):
    source, subject, inventory = staged
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    report = format_subject(subject)
    assert report["converted"] == 1 and report["failed"] == 0
    record, directory, result = descriptor(subject)
    assert len(result["tables"]) == len(XSENS_SHEETS)
    assert result["observed_marker_count"] == 5
    assert result["warnings"] == []
    assert not result["marker_selection"]["rule_applied"]
    assert result["marker_selection"]["boundary_frames"] == [1, 2, 3, 4, 5]
    assert result["sampling_rate_hz"] == 60
    assert result["recorded_at_utc"] == "2030-01-02T12:30:00+00:00"
    table = next(t for t in result["tables"] if t["source_sheet"] == "Segment Velocity")
    values = pq.read_table(directory / table["file"]).to_pydict()
    assert values["Frame"] == list(range(7))
    assert values["Right Hand x"][4] == -0.1234567890123456 * 4
    assert values["Left Hand x"][3] is None
    assert table["columns"][2]["null_count"] == 1
    assert result["all_values_roundtrip_verified"]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before == record["sha256"]
    assert record["formatter_status"] == "converted" and not record["ready_for_features"]
    schema = json.loads((Path(__file__).resolve().parents[2] / "schemas/manifest.schema.json").read_text())
    Draft202012Validator(schema).validate(json.loads((subject / "manifest.json").read_text()))
    recording_schema = json.loads((Path(__file__).resolve().parents[2] / "schemas/recording.schema.json").read_text())
    Draft202012Validator(recording_schema).validate(result)


def test_incorrect_sheet_dimensions_do_not_truncate_rows(tmp_path):
    root = tmp_path / "raw"
    source = root / "SANI/SANO_BA001/XSens/BA001_t1_task1_dx.xlsx"
    acquisition(source)
    with ZipFile(source) as book:
        parts = {name: book.read(name) for name in book.namelist()}
    for name in parts:
        if name.startswith("xl/worksheets/"):
            parts[name] = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:B2"', parts[name])
    with ZipFile(source, "w") as book:
        for name, payload in parts.items():
            book.writestr(name, payload)
    output = tmp_path / "standardized"
    write_plan(make_plan(scan(root)), output)
    subject = output / "SANI/SANO_BA001"
    assert format_subject(subject)["converted"] == 1
    _, _, result = descriptor(subject)
    assert result["observed_marker_count"] == 5
    assert all(t["row_count"] == 7 for t in result["tables"] if t["source_sheet"] not in {"General Information", "Markers"})


def test_valid_cache_is_reused_and_corrupt_output_is_rebuilt(staged):
    _, subject, _ = staged
    format_subject(subject)
    assert format_subject(subject)["cached"] == 1
    _, directory, result = descriptor(subject)
    table = next(t for t in result["tables"] if t["format"] == "parquet")
    (directory / table["file"]).write_bytes(b"corrupted output")
    assert format_subject(subject)["converted"] == 1
    assert format_subject(subject)["cached"] == 1


def test_blank_formatted_marker_rows_are_preserved_but_not_counted(staged):
    source, subject, inventory = staged
    book = openpyxl.load_workbook(source)
    book["Markers"]["A1000"].number_format = "0"
    book.save(source)
    write_plan(make_plan(scan(Path(inventory["source_root"]))), subject.parents[1])
    assert format_subject(subject)["converted"] == 1
    _, directory, result = descriptor(subject)
    markers = next(t for t in result["tables"] if t["source_sheet"] == "Markers")
    assert markers["row_count"] == 999
    assert result["observed_marker_count"] == 5
    assert markers["marker_frames"] == [1, 2, 3, 4, 5]
    assert result["warnings"] == []
    assert pq.read_table(directory / markers["file"]).num_rows == 999


def test_changed_source_is_not_formatted_using_old_metadata(staged):
    source, subject, _ = staged
    source.write_bytes(source.read_bytes() + b"change")
    report = format_subject(subject)
    assert report["failed"] == 1
    manifest = json.loads((subject / "manifest.json").read_text())
    assert manifest["records"][0]["formatter_status"] == "failed"
    assert not list(subject.rglob("*.parquet"))


def test_marker_count_mismatch_is_flagged_without_discarding_markers(staged):
    source, subject, inventory = staged
    book = openpyxl.load_workbook(source)
    book["Markers"].append([6, "Marker_6"])
    book.save(source)
    root = Path(inventory["source_root"])
    write_plan(make_plan(scan(root)), subject.parents[1])
    assert format_subject(subject)["converted"] == 1
    _, directory, result = descriptor(subject)
    assert result["observed_marker_count"] == 6
    assert "marker_count_mismatch" in result["warnings"]
    assert "temporary_first_marker_ignored" not in result["warnings"]
    assert result["marker_selection"]["boundary_frames"] == [1, 2, 3, 4, 5, 6]
    assert result["marker_selection"]["ignored_marker_frames"] == []
    assert not result["marker_selection"]["valid_for_expected_repetitions"]
    markers = next(t for t in result["tables"] if t["source_sheet"] == "Markers")
    assert pq.read_table(directory / markers["file"]).num_rows == 6
    assert markers["marker_frames"] == [1, 2, 3, 4, 5, 6]
    schema = json.loads((Path(__file__).resolve().parents[2] / "schemas/recording.schema.json").read_text())
    Draft202012Validator(schema).validate(result)


def test_cached_legacy_descriptor_gets_boundary_metadata_without_rebuilding_tables(staged):
    source, subject, inventory = staged
    book = openpyxl.load_workbook(source)
    book["Markers"].append([6, "Marker_6"])
    book.save(source)
    write_plan(make_plan(scan(Path(inventory["source_root"]))), subject.parents[1])
    format_subject(subject)
    _, directory, result = descriptor(subject)
    tables_before = {t["file"]: ((directory / t["file"]).stat().st_mtime_ns,
                                hashlib.sha256((directory / t["file"]).read_bytes()).hexdigest())
                     for t in result["tables"]}
    result["marker_selection"].update(policy_id="2026-10-04-six-markers-ignore-first-temporary-v1",
                                      rule_applied=True, temporary=True, boundary_frames=[2,3,4,5,6],
                                      effective_marker_count=5, ignored_marker_indices=[0], ignored_marker_frames=[1],
                                      valid_for_expected_repetitions=True)
    result["warnings"].append("temporary_first_marker_ignored")
    descriptor_path = directory / "recording.json"
    descriptor_path.write_text(json.dumps(result))
    report = format_subject(subject)
    assert report["cached"] == 1 and report["converted"] == 0
    assert report["results"][0]["marker_metadata_refreshed"]
    assert report["results"][0]["tables_rebuilt"] == 0
    _, _, updated = descriptor(subject)
    assert updated["marker_selection"]["boundary_frames"] == [1, 2, 3, 4, 5, 6]
    assert not updated["marker_selection"]["valid_for_expected_repetitions"]
    assert updated["warnings"] == ["marker_count_mismatch"]
    tables_after = {t["file"]: ((directory / t["file"]).stat().st_mtime_ns,
                               hashlib.sha256((directory / t["file"]).read_bytes()).hexdigest())
                    for t in updated["tables"]}
    assert tables_before == tables_after
    descriptor_before = (descriptor_path.read_bytes(), descriptor_path.stat().st_mtime_ns)
    assert not format_subject(subject)["results"][0]["marker_metadata_refreshed"]
    assert descriptor_before == (descriptor_path.read_bytes(), descriptor_path.stat().st_mtime_ns)


def test_non_numeric_signal_fails_without_creating_partial_output(staged):
    source, subject, inventory = staged
    book = openpyxl.load_workbook(source)
    book["Segment Velocity"]["B3"] = "unknown"
    book.save(source)
    write_plan(make_plan(scan(Path(inventory["source_root"]))), subject.parents[1])
    assert format_subject(subject)["failed"] == 1
    assert not list(subject.rglob("*.parquet"))
    assert not list(subject.rglob(".format-*"))


def test_flagged_identity_conflict_cannot_enter_formatter(tmp_path):
    root = tmp_path / "raw"
    acquisition(root / "PAZIENTI/FG91/Xsens/FG93_01_TASK1_T1.xlsx")
    output = tmp_path / "standardized"
    write_plan(make_plan(scan(root)), output)
    subject = output / "PAZIENTI/FG91"
    report = format_subject(subject)
    assert report["converted"] == 0
    assert not list(subject.rglob("*.parquet"))


def test_unknown_source_id_fails_before_writing(staged):
    _, subject, _ = staged
    with pytest.raises(ValueError, match="source ID"):
        format_subject(subject, source_ids=["not_a_source"])
