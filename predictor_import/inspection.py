"""Bounded container/header inspection, not numerical or clinical validation."""
from __future__ import annotations

import csv
import codecs
import io
import re
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

XSENS_SHEETS = (
    "General Information", "Markers", "Segment Velocity", "Segment Position",
    "Joint Angles ZXY", "Ergonomic Joint Angles ZXY",
)


def normalize_hemisphere(value: object) -> str | None:
    """Recognize explicit side words in clinical descriptions, never lesions by inference."""
    text = str(value).strip().casefold()
    words = re.findall(r"[a-z]+", text)
    mapping = {"right": "right", "left": "left", "dx": "right", "sx": "left",
               "destro": "right", "destra": "right", "sinistro": "left", "sinistra": "left"}
    sides = {mapping[word] for word in words if word in mapping}
    if len(sides) != 1 or not words or not (words[0] in mapping or words[-1] in mapping):
        return None
    if "?" in text or set(words) & {"not", "no", "non", "possible", "uncertain", "unknown"}:
        return None
    return next(iter(sides))


def csv_delimiter(prefix: str, expected_headers: list[str] | None = None) -> str:
    candidates = {}
    for delimiter in (",", ";", "\t"):
        try:
            candidates[delimiter] = next(csv.reader(io.StringIO(prefix), delimiter=delimiter, strict=True), [])
        except csv.Error:
            continue
    if expected_headers is not None:
        matching = [d for d, headers in candidates.items() if headers == expected_headers]
        if not matching:
            raise ValueError("CSV headers changed since inventory")
        return matching[0]
    try:
        delimiter = csv.Sniffer().sniff(prefix[:8192], delimiters=",;\t").delimiter
        if len(candidates.get(delimiter, [])) > 1 or all(len(h) <= 1 for h in candidates.values()):
            return delimiter
    except csv.Error:
        pass
    # Quoted multiline values and inconsistent late rows can confuse Sniffer.
    # The first complete header is still available for bounded inventory.
    return max(candidates, key=lambda d: len(candidates[d]), default=",")


def csv_header(stream) -> list[str]:
    # Never load an entire health log merely to inventory its columns.
    prefix = codecs.getincrementaldecoder("utf-8-sig")().decode(stream.read(65536), final=False)
    if not prefix:
        return []
    return next(csv.reader(io.StringIO(prefix), delimiter=csv_delimiter(prefix)), [])


def inspect_file(path: Path) -> dict:
    result = {"inspection_status": "not_inspected", "sheets": [], "columns": [], "archive_members": []}
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            with ZipFile(path) as archive:
                payload = archive.read("xl/workbook.xml")
                workbook = ET.fromstring(payload)
                namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                result["sheets"] = [s.attrib["name"] for s in workbook.findall("m:sheets/m:sheet", namespace)]
                result["inspection_status"] = "package_metadata_readable"
        elif suffix == ".csv":
            with path.open("rb") as stream:
                result["columns"] = csv_header(stream)
            result["inspection_status"] = "header_readable"
        elif suffix == ".zip":
            with ZipFile(path) as archive:
                counts: dict[str, int] = {}
                for info in archive.infolist():
                    counts[info.filename] = counts.get(info.filename, 0) + 1
                for info in archive.infolist():
                    if info.is_dir():
                        continue
                    member_path = PurePosixPath(info.filename)
                    safe = not (member_path.is_absolute() or ".." in member_path.parts or "\\" in info.filename
                                or (member_path.parts and ":" in member_path.parts[0]))
                    member = {
                        "path": info.filename, "size_bytes": info.file_size,
                        "compressed_bytes": info.compress_size, "crc32": f"{info.CRC:08x}",
                        "safe_path": safe, "unique_path": counts[info.filename] == 1,
                        "columns": [], "inspection_status": "listed",
                    }
                    # csv_header reads at most 64 KiB, regardless of total member
                    # size. Large BBI exports still need their real headers.
                    if safe and member["unique_path"] and member_path.suffix.lower() == ".csv":
                        try:
                            with archive.open(info) as stream:
                                member["columns"] = csv_header(stream)
                            member["inspection_status"] = "header_readable"
                        except (OSError, ValueError, UnicodeError, RuntimeError, BadZipFile, csv.Error) as error:
                            member["inspection_status"] = "inspection_error"
                            member["error"] = type(error).__name__
                    result["archive_members"].append(member)
                result["inspection_status"] = "archive_directory_readable"
    except (OSError, ValueError, KeyError, UnicodeError, ET.ParseError, RuntimeError, BadZipFile, csv.Error) as error:
        result["inspection_status"] = "inspection_error"
        result["error"] = type(error).__name__
    return result


def read_patient_metadata(path: Path, expected_id: str) -> dict:
    """Read identity, arm metadata and exact timeline labels, without demographics."""
    try:
        import openpyxl
    except ImportError:
        return {"status": "dependency_unavailable", "error": "openpyxl is required for clinical metadata"}

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        identities, hemispheres, dominant_hands, timelines = set(), set(), set(), {}
        hemisphere_evidence = []
        errors = []
        try:
            for sheet in workbook:
                rows = list(sheet.iter_rows(max_row=8, values_only=True))
                for row in rows:
                    for value in row:
                        if isinstance(value, str):
                            if re.fullmatch(r"FG\d+", value.strip(), re.IGNORECASE):
                                identities.add(value.strip().upper())
                for index, row in enumerate(rows[:-1]):
                    following = rows[index + 1]
                    for column, value in enumerate(row):
                        if not isinstance(value, str) or column >= len(following):
                            continue
                        label = value.strip().casefold()
                        if label in {"hemisphere", "dominant hand"}:
                            raw = following[column]
                            side = {"right": "right", "left": "left", "dx": "right", "sx": "left",
                                    "destro": "right", "sinistro": "left"}.get(str(raw).strip().casefold())
                            if label == "hemisphere":
                                side = normalize_hemisphere(raw)
                                hemisphere_evidence.append({"sheet": sheet.title,
                                    "cell": f"{openpyxl.utils.get_column_letter(column + 1)}{index + 2}",
                                    "source_value": str(raw) if raw is not None else None, "normalized_side": side})
                            if side:
                                (hemispheres if label == "hemisphere" else dominant_hands).add(side)
                            elif raw is not None:
                                errors.append(f"unrecognized_{label.replace(' ', '_')}")
                        if re.fullmatch(r"T\d+(?:_d\d+)?", value.strip(), re.IGNORECASE):
                            raw = following[column]
                            if hasattr(raw, "date"):
                                key, date = value.strip().upper(), raw.date().isoformat()
                                if key in timelines and timelines[key] != date:
                                    errors.append("conflicting_timeline_date")
                                timelines[key] = date
            if identities != {expected_id}:
                return {"status": "identity_conflict", "observed_ids": sorted(identities)}
            if len(hemispheres) > 1 or len(dominant_hands) > 1 or errors:
                return {"status": "metadata_conflict", "errors": sorted(set(errors))}
            return {"status": "available", "lesion_hemisphere": next(iter(hemispheres), None),
                    "dominant_hand": next(iter(dominant_hands), None), "timeline_dates": timelines,
                    "hemisphere_evidence": hemisphere_evidence}
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError, ET.ParseError, RuntimeError, BadZipFile) as error:
        return {"status": "read_error", "error": type(error).__name__}
