"""Xsens XLSX to Parquet/JSON with complete value round-trip verification."""
from __future__ import annotations

from datetime import date, datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import shutil
import tempfile
import time
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from . import SCHEMA_VERSION
from .adapter import atomic_json
from .inventory import now
from .markers import select_marker_boundaries
from .storage_names import xsens_table_names, xsens_targets

FORMATTER_VERSION = "xsens-parquet-0.2.1"
CHUNK_ROWS = 1024


def checksum(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(chunk)
    return result.hexdigest()


def fingerprint(path: Path) -> tuple[int, int]:
    stat = path.stat()
    return stat.st_size, stat.st_mtime_ns


def feed_rows(digest, rows) -> None:
    for row in rows:
        digest.update(json.dumps(row, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode("utf-8"))
        digest.update(b"\n")


def json_value(value):
    if isinstance(value, (datetime, date)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Non-finite metadata value cannot be serialized losslessly")
    return value


def metadata_sheet(sheet, output: Path) -> dict:
    # Some exports declare dimensions smaller than their actual XML data.
    sheet.reset_dimensions()
    rows = []
    for row in sheet.iter_rows():
        rows.append([{"data_type": cell.data_type, "value": json_value(cell.value)} for cell in row])
    document = {"source_sheet": sheet.title, "rows": rows}
    atomic_json(output, document)
    if json.loads(output.read_text(encoding="utf-8")) != document:
        raise ValueError("Metadata JSON round-trip mismatch")
    return {"source_sheet": sheet.title, "file": output.name, "format": "json", "row_count": len(rows),
            "sha256": checksum(output), "roundtrip_verified": True}


def schema_for(headers: list[str], rows: list[tuple], *, markers: bool) -> pa.Schema:
    if not headers or any(not isinstance(h, str) or not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError("Empty, duplicate or non-string worksheet headers")
    fields = []
    for column, name in enumerate(headers):
        if name == "Frame":
            dtype = pa.int64()
        elif not markers:
            dtype = pa.float64()
        else:
            values = [r[column] for r in rows if column < len(r) and r[column] is not None]
            if not values:
                dtype = pa.null()
            elif all(isinstance(v, str) for v in values):
                dtype = pa.string()
            elif all(isinstance(v, bool) for v in values):
                dtype = pa.bool_()
            elif all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in values):
                dtype = pa.float64()
            else:
                raise ValueError(f"Unsupported mixed marker column {name!r}; original remains preserved")
        fields.append(pa.field(name, dtype))
    return pa.schema(fields)


def normalize_row(row, schema: pa.Schema) -> list:
    if len(row) > len(schema):
        raise ValueError("Worksheet data extends beyond its headers")
    values = list(row) + [None] * (len(schema) - len(row))
    result = []
    for field, value in zip(schema, values):
        if value is None:
            result.append(None)
        elif pa.types.is_integer(field.type):
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or int(value) != value or not -(2**63) <= value < 2**63:
                raise ValueError(f"Non-integral/out-of-range Frame value {value!r}")
            result.append(int(value))
        elif pa.types.is_floating(field.type):
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                raise ValueError(f"Non-numeric/non-finite value in {field.name!r}")
            if isinstance(value, int) and abs(value) > 2**53:
                raise ValueError("Integer would lose precision in float64")
            result.append(float(value))
        elif pa.types.is_string(field.type) and isinstance(value, str):
            result.append(value)
        elif pa.types.is_boolean(field.type) and isinstance(value, bool):
            result.append(value)
        else:
            raise ValueError(f"Unexpected value/type in {field.name!r}")
    return result


def table_sheet(sheet, output: Path, *, markers: bool) -> dict:
    sheet.reset_dimensions()
    iterator = sheet.iter_rows(values_only=True)
    headers = list(next(iterator, ()))
    # Markers are small and may contain text; infer their types from all rows.
    marker_rows = list(iterator) if markers else None
    schema = schema_for(headers, marker_rows or [], markers=markers)
    rows = iter(marker_rows) if markers else iterator
    source_digest, count, nulls = hashlib.sha256(), 0, [0] * len(schema)
    first_frame = last_frame = None
    frame_digest = hashlib.sha256()
    marker_frames = []
    pending = []
    writer = pq.ParquetWriter(output, schema, compression="zstd")
    try:
        for row in rows:
            values = normalize_row(row, schema)
            feed_rows(source_digest, [values])
            count += 1
            for i, value in enumerate(values):
                nulls[i] += value is None
            if "Frame" in headers:
                frame = values[headers.index("Frame")]
                feed_rows(frame_digest, [[frame]])
                if count == 1:
                    first_frame = frame
                last_frame = frame
                # Style-only worksheet rows are preserved in Parquet but are
                # not marker events. A nonempty event with no frame stays
                # visible as None and triggers the marker validity warning.
                if markers and any(value is not None for value in values):
                    marker_frames.append(frame)
            pending.append(values)
            if len(pending) == CHUNK_ROWS:
                writer.write_table(pa.Table.from_arrays([pa.array([r[i] for r in pending], type=f.type) for i, f in enumerate(schema)], schema=schema))
                pending.clear()
        if pending:
            writer.write_table(pa.Table.from_arrays([pa.array([r[i] for r in pending], type=f.type) for i, f in enumerate(schema)], schema=schema))
    finally:
        writer.close()
    # Verify EVERY value in original order, not just counts or a few samples.
    parquet = pq.ParquetFile(output)
    output_digest = hashlib.sha256()
    for batch in parquet.iter_batches(batch_size=CHUNK_ROWS):
        columns = [column.to_pylist() for column in batch.columns]
        feed_rows(output_digest, zip(*columns))
    if parquet.metadata.num_rows != count or output_digest.digest() != source_digest.digest():
        raise ValueError(f"Parquet round-trip mismatch in {sheet.title}")
    return {"source_sheet": sheet.title, "file": output.name, "format": "parquet", "row_count": count,
            "columns": [{"name": f.name, "dtype": str(f.type), "unit": None, "null_count": n} for f, n in zip(schema, nulls)],
            "header_row": 1, "first_frame": first_frame, "last_frame": last_frame,
            "frame_values_sha256": frame_digest.hexdigest() if "Frame" in headers else None,
            "marker_frames": marker_frames if markers else None,
            "values_sha256": source_digest.hexdigest(), "sha256": checksum(output), "roundtrip_verified": True}


def cached_descriptor(target: Path, source_sha: str, policy_id: str) -> dict | None:
    try:
        descriptor = json.loads((target / "recording.json").read_text(encoding="utf-8"))
        if descriptor["source_sha256"] != source_sha or descriptor["formatter_version"] != FORMATTER_VERSION or descriptor["import_policy_id"] != policy_id:
            return None
        for table in descriptor["tables"]:
            file = target / table["file"]
            if not file.resolve().is_relative_to(target.resolve()) or checksum(file) != table["sha256"]:
                return None
        return descriptor
    except (OSError, ValueError, KeyError):
        return None


def refresh_marker_selection(descriptor: dict) -> bool:
    """Refresh derived boundary metadata while retaining source diagnostics."""
    frames = next(t["marker_frames"] for t in descriptor["tables"]
                  if t["source_sheet"] == "Markers")
    selection = select_marker_boundaries(
        frames, expected_marker_count=descriptor["expected_marker_count"],
        expected_repetitions=descriptor["expected_repetitions"],
    )
    warnings = [w for w in descriptor["warnings"]
                if w != "temporary_first_marker_ignored"]
    if selection["rule_applied"]:
        warnings.append("temporary_first_marker_ignored")
    changed = (descriptor.get("marker_selection") != selection
               or descriptor["warnings"] != warnings)
    descriptor["marker_selection"] = selection
    descriptor["warnings"] = warnings
    return changed


def convert_record(source: Path, record: dict, target: Path, policy_id: str) -> dict:
    start = time.perf_counter()
    if fingerprint(source) != (record["size_bytes"], record["mtime_ns"]):
        raise ValueError("Source changed since inventory; refresh before converting")
    try:
        previous = json.loads((target / "recording.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        previous = {}
    if previous.get("source_id") and previous["source_id"] != record["source_id"]:
        raise ValueError("Xsens destination belongs to a different source; use a distinct acquisition name")
    before = fingerprint(source)
    source_sha = checksum(source)
    cache = cached_descriptor(target, source_sha, policy_id) if target.exists() else None
    if cache and (cache.get("source_id") != record["source_id"] or cache.get("source_path") != record["source_path"]):
        raise ValueError("Cached Xsens descriptor identity disagrees with source")
    if cache and fingerprint(source) == before:
        refreshed = refresh_marker_selection(cache)
        if refreshed:
            atomic_json(target / "recording.json", cache)
        return {"status": "cached", "source_id": record["source_id"], "descriptor": cache,
                "marker_metadata_refreshed": refreshed,
                "tables_reused": len(cache["tables"]), "tables_rebuilt": 0}
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix=".format-", dir=target.parent))
    try:
        with ZipFile(source) as archive:
            damaged = archive.testzip()
            if damaged:
                raise ValueError(f"XLSX CRC error in {damaged}")
        book = openpyxl.load_workbook(source, read_only=True, data_only=False)
        tables, metadata = [], {}
        try:
            filenames = xsens_table_names(book.sheetnames)
            for sheet in book:
                if sheet.title == "General Information":
                    tables.append(metadata_sheet(sheet, temporary / "general_information.json"))
                    rows = json.loads((temporary / "general_information.json").read_text())["rows"]
                    metadata = {str(row[0]["value"]): row[1]["value"] for row in rows if len(row) >= 2 and row[0]["value"] is not None}
                else:
                    tables.append(table_sheet(sheet, temporary / filenames[sheet.title], markers=sheet.title == "Markers"))
        finally:
            book.close()
        final_checksum = checksum(source)
        if fingerprint(source) != before or final_checksum != source_sha:
            raise ValueError("Source changed during conversion")
        markers = next(t for t in tables if t["source_sheet"] == "Markers")
        marker_frames = markers["marker_frames"]
        warnings = []
        if record["expected_marker_count"] is not None and len(marker_frames) != record["expected_marker_count"]:
            warnings.append("marker_count_mismatch")
        if any(v is None for v in marker_frames) or any(a >= b for a, b in zip(marker_frames, marker_frames[1:]) if a is not None and b is not None):
            warnings.append("markers_not_strictly_increasing")
        signals = [t for t in tables if t["format"] == "parquet" and t["source_sheet"] != "Markers"]
        if len({t["frame_values_sha256"] for t in signals}) > 1:
            warnings.append("signal_frame_sequences_differ")
        utc = metadata.get("Recorded Date UTC")
        recorded_at = None
        if isinstance(utc, dict) and utc.get("type") == "datetime":
            dt = datetime.fromisoformat(utc["value"])
            recorded_at = dt.replace(tzinfo=timezone.utc).isoformat() if dt.tzinfo is None else dt.astimezone(timezone.utc).isoformat()
        descriptor = {"schema_version": SCHEMA_VERSION, "formatter_version": FORMATTER_VERSION,
            "import_policy_id": policy_id, "source_id": record["source_id"], "source_path": record["source_path"],
            "source_sha256": source_sha, "source_size_bytes": before[0], "source_mtime_ns": before[1],
            "created_at": now(), "compression": "zstd", "tables": tables, "recorded_at_utc": recorded_at,
            "sampling_rate_hz": metadata.get("Frame Rate"), "unit_policy": "unknown_unless_declared_by_source",
            "expected_repetitions": record["expected_repetitions"], "expected_marker_count": record["expected_marker_count"],
            "observed_marker_count": len(marker_frames), "warnings": warnings,
            "all_values_roundtrip_verified": True, "ready_for_features": False,
            "output_size_bytes": sum((temporary / table["file"]).stat().st_size for table in tables),
            "conversion_seconds": time.perf_counter() - start}
        refresh_marker_selection(descriptor)
        atomic_json(temporary / "recording.json", descriptor)
        # Build a complete new directory before replacing a prior conversion.
        backup = target.with_name(target.name + ".previous")
        if backup.exists():
            raise ValueError("Unresolved previous conversion backup; inspect before replacing")
        if target.exists():
            target.rename(backup)
        try:
            temporary.rename(target)
        except OSError:
            if backup.exists():
                backup.rename(target)
            raise
        if backup.exists():
            shutil.rmtree(backup)
        return {"status": "converted", "source_id": record["source_id"], "descriptor": descriptor,
                "tables_reused": 0, "tables_rebuilt": len(tables)}
    finally:
        if temporary.exists():
            shutil.rmtree(temporary)


def format_subject(subject_path: Path, *, source_ids: list[str] | None = None, limit: int | None = None) -> dict:
    subject_path = subject_path.resolve()
    manifest_path = subject_path / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("schema_version") != SCHEMA_VERSION:
        raise ValueError("Restage metadata with the current schema before formatting")
    root = Path(manifest["source_root"]).resolve()
    if subject_path == root or subject_path.is_relative_to(root) or root.is_relative_to(subject_path):
        raise ValueError("Formatted output must be separate from raw sources")
    known_ids = {r["source_id"] for r in manifest["records"]}
    if source_ids and set(source_ids) - known_ids:
        raise ValueError("Requested source ID not found in subject manifest")
    from .layout import migrate_subject
    targets = xsens_targets(manifest["records"])
    selected = [r for r in manifest["records"] if r["source_id"] in targets
                and (not source_ids or r["source_id"] in source_ids)]
    if limit is not None:
        selected = selected[:limit]
    migration = migrate_subject(subject_path, targets={r["source_id"]: targets[r["source_id"]] for r in selected})
    if migration["records_changed"]:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    results, count = [], 0
    for record in manifest["records"]:
        if record["disposition"] != "mapped" or record["artifact_kind"] != "xsens_workbook" or (source_ids and record["source_id"] not in source_ids):
            continue
        if limit is not None and count >= limit:
            break
        count += 1
        try:
            if not record["normalized_target"]:
                raise ValueError("Recording has no canonical destination")
            target = subject_path / record["normalized_target"]
            source = root / record["source_path"]
            if target.resolve() == subject_path or not target.resolve().is_relative_to(subject_path) or source.is_symlink() or not source.resolve().is_relative_to(root):
                raise ValueError("Source or destination escapes its root")
            result = convert_record(source, record, target, manifest["import_policy"]["policy_id"])
            record["sha256"] = result["descriptor"]["source_sha256"]
            record["formatter_status"] = "converted"
            record["formatter_report"] = {"descriptor_path": record["normalized_target"] + "/recording.json", "warnings": result["descriptor"]["warnings"]}
            result = {k: v for k, v in result.items() if k != "descriptor"}
        except (OSError, ValueError, KeyError, pa.ArrowException, BadZipFile, ParseError) as error:
            record["formatter_status"] = "failed"
            record["formatter_report"] = {"error": str(error)}
            result = {"status": "failed", "source_id": record["source_id"], "error": str(error)}
        results.append(result)
        # Preserve completed work even if a later file fails or the run is interrupted.
        manifest["adapter_stage"] = "metadata_and_tables"
        atomic_json(manifest_path, manifest)
    report = {"schema_version": SCHEMA_VERSION, "participant_key": manifest["participant_key"],
              "results": results, "converted": sum(r["status"] == "converted" for r in results),
              "cached": sum(r["status"] == "cached" for r in results), "failed": sum(r["status"] == "failed" for r in results),
              "ready_for_feature_extraction": False}
    report["layout_migration"] = migration
    atomic_json(subject_path / "format_report.json", report)
    return report
