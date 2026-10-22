"""Portable, verified patient XLSX export from current Xsens feature tables."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
from pathlib import Path
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font

from predictor_import.markers import MARKER_POLICY_ID
from shared_code.reader import DataContractError
from .xsens_features import FEATURE_NAMES, METHOD_ID

HEADERS = ("subject", "task", "repetition (1-4)", "arm (dx-sx)",
           "healthy/paretic", "number of samples", *FEATURE_NAMES)


def patient_destination(output: Path, participant_key: str) -> Path:
    parts = participant_key.split("/")
    if len(parts) != 2 or parts[0] != "PAZIENTI" or parts[1] in {"", ".", ".."} or "\\" in parts[1]:
        raise DataContractError("export_identity_invalid", "Patient XLSX export requires PAZIENTI/<subject>")
    root = output.resolve()
    destination = root / "PAZIENTI" / parts[1] / "Xsens_tasks" / "Xsens_tasks.xlsx"
    if not destination.resolve().is_relative_to(root):
        raise DataContractError("unsafe_output", "Workbook destination escapes output root")
    return destination


def _read_rows(directory: Path):
    metadata = json.loads((directory / "metadata.json").read_text())
    method = metadata["method"]
    if method["method_id"] != METHOD_ID or method["rom"]["marker_policy_id"] != MARKER_POLICY_ID:
        raise DataContractError("export_method_mismatch", "Export requires the current quaternion elbow method")
    source = directory / "features.csv"
    payload = source.read_bytes()
    if hashlib.sha256(payload).hexdigest() != metadata["artifacts"]["features.csv"]["sha256"]:
        raise DataContractError("export_checksum_mismatch", "Feature CSV checksum differs from extraction metadata")
    groups, seen = {}, set()
    with io.StringIO(payload.decode("utf-8"), newline="") as stream:
        reader = csv.DictReader(stream)
        required = {"participant_key", "timepoint", "task", "repetition", "anatomical_side", "limb_condition",
                    "n_samples", "method_id", "method_sha256", "marker_policy_id", *FEATURE_NAMES}
        if not required <= set(reader.fieldnames or []):
            raise DataContractError("export_columns_missing", "Feature CSV does not have the required columns")
        for row in reader:
            if (row["method_id"] != METHOD_ID or row["method_sha256"] != method["method_sha256"]
                    or row["marker_policy_id"] != MARKER_POLICY_ID):
                raise DataContractError("export_method_mismatch", "Feature rows contain mixed methods")
            participant = row["participant_key"]
            patient_destination(Path("."), participant)  # Validate the cohort-qualified identity.
            visit = row["timepoint"] or "unlabelled"
            if len(visit) > 31 or any(c in visit for c in "\\/*?:[]") or visit.startswith("'") or visit.endswith("'"):
                raise DataContractError("export_sheet_invalid", "Unsupported visit worksheet name")
            task, repetition, samples = (int(row[name]) for name in ("task", "repetition", "n_samples"))
            arm = {"right": "dx", "left": "sx"}.get(row["anatomical_side"])
            condition = {"non_paretic": "healthy", "paretic": "paretic"}.get(row["limb_condition"])
            if task not in {1, 2, 3, 4} or repetition not in {1, 2, 3, 4} or samples <= 0 or arm is None or condition is None:
                raise DataContractError("export_row_invalid", "Invalid task, repetition, side, condition or sample count")
            key = (participant, visit, task, repetition, arm, condition)
            if key in seen:
                raise DataContractError("export_duplicate_row", "Ambiguous repeated acquisition in a visit worksheet")
            seen.add(key)
            values = [float(row[name]) if row[name] != "" else None for name in FEATURE_NAMES]
            if any(value is not None and not math.isfinite(value) for value in values):
                raise DataContractError("export_nonfinite_value", "Feature table contains a populated nonfinite value")
            groups.setdefault(participant, {}).setdefault(visit, []).append(
                [participant.split("/")[1], task, repetition, arm, condition, samples, *values])
    if not seen:
        raise DataContractError("export_empty", "No patient task feature rows to export")
    for visits in groups.values():
        # Excel worksheet names are case-insensitive.
        if len({name.casefold() for name in visits}) != len(visits):
            raise DataContractError("export_sheet_collision", "Visit names collide in Excel")
        for rows in visits.values():
            rows.sort(key=lambda r: (r[1], r[3], r[2]))
    return groups


def _verify(path: Path, visits: dict):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if book.sheetnames != list(visits):
            raise DataContractError("export_roundtrip_failed", "Workbook worksheet order changed")
        for name, expected_rows in visits.items():
            actual = list(book[name].values)
            if tuple(actual[0]) != HEADERS or len(actual) != len(expected_rows) + 1:
                raise DataContractError("export_roundtrip_failed", "Workbook header or row count differs")
            for actual_row, expected in zip(actual[1:], expected_rows):
                if list(actual_row[:6]) != expected[:6]:
                    raise DataContractError("export_roundtrip_failed", "Workbook row identity differs")
                for got, value in zip(actual_row[6:], expected[6:]):
                    equal = got is None if value is None else isinstance(got, (int, float)) and math.isclose(got, value, rel_tol=1e-13, abs_tol=1e-10)
                    if not equal:
                        raise DataContractError("export_roundtrip_failed", "Workbook feature value differs")
    finally:
        book.close()


def export_workbooks(feature_directories: list[Path], output: Path) -> list[dict]:
    """Preflight every destination, verify saved values, and publish without replacement."""
    groups = {}
    for directory in feature_directories:
        for participant, visits in _read_rows(directory).items():
            if participant in groups:
                raise DataContractError("export_duplicate_subject", "Supply one feature directory per patient")
            groups[participant] = visits
    destinations = {key: patient_destination(output, key) for key in groups}
    for destination in destinations.values():
        if destination.exists():
            raise DataContractError("output_exists", f"Workbook already exists: {destination}")
    outputs = []
    for participant, visits in groups.items():
        visits = dict(sorted(visits.items(), key=lambda item: [int(p) if p.isdigit() else p for p in re.split(r"(\d+)", item[0])]))
        destination = destinations[participant]
        destination.parent.mkdir(parents=True, exist_ok=True)
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, rows in visits.items():
            sheet = book.create_sheet(name)
            sheet.append(HEADERS)
            for row in rows:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = True
            sheet.row_dimensions[1].height = 32
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for column in range(1, len(HEADERS) + 1):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 18 if 2 <= column <= 6 else 24
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    elif cell.column in {2, 3, 6}:
                        cell.number_format = "0"
                    elif cell.column >= 7:
                        cell.number_format = "0.000000"
        fd, temporary = tempfile.mkstemp(prefix=".xsens-", suffix=".xlsx", dir=destination.parent)
        os.close(fd)
        try:
            book.save(temporary)
            _verify(Path(temporary), visits)
            os.link(temporary, destination)  # Atomic and exclusive: never replaces an existing workbook.
        finally:
            book.close()
            Path(temporary).unlink(missing_ok=True)
        outputs.append({"participant_key": participant, "path": str(destination),
                        "rows": sum(map(len, visits.values())),
                        "sha256": hashlib.sha256(destination.read_bytes()).hexdigest()})
    return outputs
