from __future__ import annotations

import ast
import hashlib
import json
from pathlib import Path

import numpy as np
import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from scipy.signal import butter, sosfiltfilt

from predictor_import.adapter import atomic_json, make_plan, write_plan
from predictor_import.formatter import format_subject
from predictor_import.inspection import XSENS_SHEETS
from predictor_import.inventory import scan
from xsens_code.__main__ import main
from shared_code.reader import DataContractError, SubjectReader
from xsens_code.rom import ROM_COLUMNS
from xsens_code.smoothness import spectral_arc_length
from xsens_code.xsens_cli import extract_subject
from xsens_code.xsens_dtw import (DTW_COORDINATE_METHOD, ORIENTATION_SHEET, align_and_mirror,
                                 center_trajectory, hand_distance, initial_heading_frame, pair_hand_dtw)
from xsens_code.xsens_features import COMP_COLUMNS, FEATURE_NAMES, SEGMENTS, XsensConfig, extract_xsens
from xsens_code.tests.test_xsens_reader import opened, replace_column, rewrite_table


@pytest.fixture
def full_subject(tmp_path):
    raw = tmp_path / "raw"
    for file_side in ("dx", "sx"):
        source = raw / f"SANI/SANO_FG91/XSens/FG91_t1_task1_{file_side}.xlsx"
        source.parent.mkdir(parents=True, exist_ok=True)
        book = openpyxl.Workbook()
        book.remove(book.active)
        t = np.arange(240) / 60
        variation = 1 if file_side == "dx" else 1.2
        for name in (*XSENS_SHEETS, ORIENTATION_SHEET):
            sheet = book.create_sheet(name)
            if name == "General Information":
                sheet.append(["Frame Rate", 60])
                continue
            if name == "Markers":
                sheet.append(["Frame", "Marker Text"])
                for i, frame in enumerate([1010, 1060, 1110, 1160, 1210]):
                    sheet.append([frame, f"event{i}"])
                continue
            if name == ORIENTATION_SHEET:
                columns = [f"T8 q{i}" for i in range(4)]
                values = [np.ones_like(t), np.zeros_like(t), np.zeros_like(t), np.zeros_like(t)]
                bend = np.deg2rad(40 + 20 * np.sin(2 * np.pi * .7 * t))
                for side in ("Right", "Left"):
                    for segment in ("Upper Arm", "Forearm"):
                        columns += [f"{side} {segment} q{i}" for i in range(4)]
                        values += ([np.ones_like(t), np.zeros_like(t), np.zeros_like(t), np.zeros_like(t)]
                            if segment == "Upper Arm" else [np.cos(bend/2), np.zeros_like(t), np.zeros_like(t), np.sin(bend/2)])
            elif name == "Joint Angles ZXY":
                columns = [f"{side} {suffix}" for side in ("Right", "Left") for suffix in [*ROM_COLUMNS.values(), "Elbow Flexion/Extension"]]
                values = [30 + (10 + j) * np.sin(2 * np.pi * 0.7 * t) + 2 * np.sin(2 * np.pi * 12 * t) for j in range(len(columns))]
                columns += [f"{side} Elbow Ulnar Deviation/Radial Deviation" for side in ("Right", "Left")]
                values += [np.zeros_like(t), np.zeros_like(t)]
            elif name == "Segment Velocity":
                bases = ["T8", *[f"{side} {segment}" for side in ("Right", "Left") for segment in list(SEGMENTS.values())[1:]]]
                columns = [f"{base} {axis}" for base in bases for axis in "xyz"]
                values = [variation * (0.2 + (j + 1) / 100 + 0.06 * np.sin(2 * np.pi * (0.7 + j / 20) * t)
                                      + 0.08 * np.sin(2 * np.pi * 17 * t)) for j in range(len(columns))]
            elif name == "Segment Position":
                columns = [f"{side} Hand {axis}" for side in ("Right", "Left") for axis in "xyz"]
                values = [variation * 0.1 * np.sin(2 * np.pi * (0.4 + j / 10) * t) + j / 3 for j in range(len(columns))]
                bend = np.deg2rad(40 + 20 * np.sin(2 * np.pi * .7 * t))
                for side, direction, offset in (("Right", -1, 0), ("Left", 1, 3)):
                    hand = np.column_stack(values[offset:offset+3])
                    upper_axis = np.column_stack([np.zeros_like(t), np.ones_like(t), np.zeros_like(t)])
                    forearm_axis = np.column_stack([-np.sin(bend), np.cos(bend), np.zeros_like(t)])
                    forearm = hand - direction * forearm_axis * .3
                    upper = forearm - direction * upper_axis * .3
                    for segment, points in (("Upper Arm", upper), ("Forearm", forearm)):
                        columns += [f"{side} {segment} {axis}" for axis in "xyz"]
                        values += [points[:,j] for j in range(3)]
            elif name == "Ergonomic Joint Angles ZXY":
                columns = list(COMP_COLUMNS.values())
                values = [3 + 5 * np.sin(2 * np.pi * 0.6 * t), 5 + 10 * np.sin(2 * np.pi * 0.7 * t)]
            else:
                columns, values = ["unused"], [t]
            sheet.append(["Frame", *columns])
            for i in range(240):
                sheet.append([1000 + i, *[float(v[i]) for v in values]])
        book.save(source)
    output = tmp_path / "standardized"
    write_plan(make_plan(scan(raw)), output)
    subject = output / "SANI/SANO_FG91"
    assert format_subject(subject)["converted"] == 2
    return subject


def legacy_oracle():
    """Execute actual archived numerical code in a test-only namespace."""
    root = Path(__file__).resolve().parents[2] / "archive/xsens_code/src"
    context = {"FS_HZ": 60., "__name__": __name__}
    for name in ("preprocessing.py", "sparc.py", "features.py"):
        tree = ast.parse((root / name).read_text())
        nodes = [node for node in tree.body if not (isinstance(node, ast.ImportFrom) and
                                                  (node.level or node.module == "__future__"))]
        exec(compile(ast.Module(body=nodes, type_ignores=[]), str(root / name), "exec"), context)
    return context


def independent_dtw(a, b):
    """Full cost matrix, independent of dtaidistance and production helpers."""
    a = a - (a[0] if len(a) < 10 else a[:10].mean(axis=0))
    b = b - (b[0] if len(b) < 10 else b[:10].mean(axis=0))
    matrix = np.full((len(a) + 1, len(b) + 1), np.inf)
    matrix[0, 0] = 0
    for i in range(1, len(a) + 1):
        for j in range(1, len(b) + 1):
            matrix[i, j] = np.dot(a[i - 1] - b[j - 1], a[i - 1] - b[j - 1]) + min(
                matrix[i - 1, j], matrix[i, j - 1], matrix[i - 1, j - 1])
    return float(np.sqrt(matrix[-1, -1]))


def test_original_angles_and_independent_vector_speed_and_mirrored_dtw(full_subject):
    legacy = legacy_oracle()
    reader = SubjectReader(full_subject)
    report = extract_subject(reader, XsensConfig())
    pre = {}
    mapping = {"segment_velocity": "Segment Velocity", "segment_position": "Segment Position",
               "joint_angles_zxy": "Joint Angles ZXY", "ergo_joint_angles_zxy": "Ergonomic Joint Angles ZXY"}
    for result in report["results"]:
        record = next(r for r in reader.manifest["records"] if r["source_id"] == result["source_id"])
        recording = reader.open_xsens(record)
        side = "dx" if record["anatomical_side"] == "right" else "sx"
        sheets = {k: recording.read_table(v).to_pandas() for k, v in mapping.items()}
        windows = [(r["row_start"], r["row_end_exclusive"]) for r in result["feature_rows"]]
        pre[side] = legacy["preprocess_file"](sheets, {"arto": side}, windows)
        expected = legacy["extract_features"](pre[side])
        for actual, old in zip(result["feature_rows"], expected):
            for f in (*ROM_COLUMNS, *COMP_COLUMNS):
                assert actual[f] == old[f]
            for role, segment in SEGMENTS.items():
                base = segment if role == "trunk" else f"{record['anatomical_side'].title()} {segment}"
                xyz = sheets['segment_velocity'][[f"{base} {axis}" for axis in "xyz"]].to_numpy()
                for prefix, cutoff in (("vel_mean", 6), ("sparc", 15)):
                    vector = sosfiltfilt(butter(4, cutoff, fs=60, output="sos"), xyz, axis=0, padlen=12)
                    speed = np.linalg.norm(vector, axis=1)[actual['row_start']:actual['row_end_exclusive']]
                    value = speed.mean() if prefix == "vel_mean" else legacy['sparc'](speed, 60)[0]
                    assert actual[f"{prefix}_{role}"] == pytest.approx(value, abs=1e-9)
            assert actual["n_samples"] == 50 and actual["qc_pass_feature_count"] == 15
        assert set(result["provenance"]["checked_tables"]) == {"General Information", "Markers", ORIENTATION_SHEET, *mapping.values()}
    for result in report["results"]:
        for i, row in enumerate(result["feature_rows"]):
            left = pre["sx"].hand_position_6hz[i] * [1, -1, 1]
            assert row["dtw_hand"] == pytest.approx(independent_dtw(pre["dx"].hand_position_6hz[i], left), abs=1e-12)
    assert report["computed_feature_values"] == report["qc_pass_feature_values"] == 120
    assert len(report["pairs"]) == 4


@pytest.mark.parametrize("lengths", [(12, 17), (4, 5), (10, 10)])
def test_dtw_centering_is_first_ten_not_centroid_and_translation_invariant(lengths):
    rng = np.random.default_rng(16)
    a, b = [rng.normal(size=(n, 3)).cumsum(axis=0) for n in lengths]
    assert hand_distance(a, b) == pytest.approx(independent_dtw(a, b), abs=1e-12)
    assert hand_distance(a + 500, b - 20) == pytest.approx(hand_distance(a, b), abs=1e-10)
    assert np.allclose(center_trajectory(a), a - (a[0] if len(a) < 10 else a[:10].mean(axis=0)))
    if len(a) > 10:
        assert not np.allclose(center_trajectory(a), a - a.mean(axis=0))


@pytest.mark.parametrize("invalid", [None, float("nan"), float("inf")])
def test_bad_velocity_outside_windows_excludes_only_segment_mean_and_sparc(full_subject, invalid):
    column = "Right Shoulder x"
    def change(table):
        values = table[column].to_pylist()
        values[0] = invalid
        return replace_column(table, column, values)
    rewrite_table(full_subject, "Segment Velocity", change)
    result = extract_xsens(opened(full_subject)).report
    assert result["computed_feature_values"] == 48  # 12 usable groups; DTW awaits a pair
    for row in result["feature_rows"]:
        assert row["vel_mean_shoulder"] is None and row["sparc_shoulder"] is None
        assert all(row[f] is not None for f in FEATURE_NAMES if f not in {"vel_mean_shoulder", "sparc_shoulder", "dtw_hand"})
    issue = next(f for f in result["channels"]["vel_mean_shoulder"]["findings"] if f["severity"] == "error")
    assert issue["first_frame"] == 1000 and issue["first_source_excel_row"] == 2


def test_bad_hand_keeps_other_features_and_excludes_both_paired_distances(full_subject):
    rewrite_table(full_subject, "Segment Position", lambda t: replace_column(t, "Right Hand x", [None] * 240))
    report = extract_subject(SubjectReader(full_subject), XsensConfig())
    assert report["computed_feature_values"] == 112
    assert all(row["dtw_hand"] is None for r in report["results"] for row in r["feature_rows"])
    assert not report["pairs"]


def test_compensation_unit_conflict_and_jump_have_local_scope(full_subject):
    recording = opened(full_subject)
    desc = recording.table_descriptor("Ergonomic Joint Angles ZXY")
    desc["columns"][1]["unit"] = "rad"
    atomic_json(recording.descriptor_path, recording.descriptor)
    column = "Vertical_T8 Flexion/Extension"
    def change(table):
        values = ((np.linspace(170, 200, 240) + 180) % 360 - 180).tolist()
        return replace_column(table, column, values)
    rewrite_table(full_subject, "Ergonomic Joint Angles ZXY", change)
    result = extract_xsens(opened(full_subject)).report
    for row in result["feature_rows"]:
        assert row["comp_trunk_lateral_bending"] is None
        assert row["feature_status"]["comp_trunk_flex_ext"] == "computed"
        assert row["comp_trunk_flex_ext"] is not None
        assert row["feature_status"]["rom_elbow_bending"] == "computed"
        expected = sosfiltfilt(butter(4, 6, fs=60, output="sos"), np.linspace(170, 200, 240), padlen=12)
        assert row["comp_trunk_flex_ext"] == pytest.approx(np.ptp(expected[row["row_start"]:row["row_end_exclusive"]]), abs=1e-9)
        assert "before_unwrap" in row["feature_evidence"]["comp_trunk_flex_ext"]


def test_zero_speed_mean_is_zero_but_sparc_is_explicitly_null(full_subject):
    def change(table):
        for axis in "xyz":
            table = replace_column(table, f"T8 {axis}", [0.] * 240)
        return table
    rewrite_table(full_subject, "Segment Velocity", change)
    result = extract_xsens(opened(full_subject)).report
    assert all(r["vel_mean_trunk"] == 0 and r["sparc_trunk"] is None for r in result["feature_rows"])
    assert any(f["code"] == "sparc_degenerate_spectrum" for f in result["channels"]["sparc_trunk"]["findings"])
    with pytest.raises(DataContractError, match="two samples"):
        spectral_arc_length(np.asarray([1.]), 60)
    with pytest.raises(DataContractError):
        spectral_arc_length(np.zeros(20), 60)


def test_one_selected_spectral_bin_is_excluded_instead_of_zero_arc_length():
    with pytest.raises(DataContractError) as exc:
        spectral_arc_length(np.ones(32), 100000)
    assert exc.value.code == "sparc_degenerate_band"


def test_source_rate_controls_new_speed_filters(full_subject):
    recording = opened(full_subject)
    raw_root = Path(recording.subject.manifest["source_root"])
    source = raw_root / recording.record["source_path"]
    book = openpyxl.load_workbook(source)
    book["General Information"].cell(1, 2).value = 120
    book.save(source)
    write_plan(make_plan(scan(raw_root)), full_subject.parents[1])
    format_subject(full_subject)
    recording = opened(full_subject)
    result = extract_xsens(recording).report
    table = recording.read_table("Segment Velocity")
    xyz = np.column_stack([table[f"T8 {axis}"].to_pylist() for axis in "xyz"])
    filtered = np.linalg.norm(sosfiltfilt(butter(4, 6, fs=120, output="sos"), xyz, axis=0, padlen=12), axis=1)
    row = result["feature_rows"][0]
    assert row["vel_mean_trunk"] == pytest.approx(filtered[10:60].mean(), abs=1e-10)
    assert row["duration_s"] == 50 / 120
    paired = extract_subject(SubjectReader(full_subject), XsensConfig())
    assert all(row["dtw_hand"] is None for r in paired["results"] for row in r["feature_rows"])
    assert any(f["code"] == "dtw_pair_sampling_rate_mismatch" for r in paired["results"] for f in r["flags"])


def pair_result(source, side, condition="non_paretic", **metadata):
    row = {"repetition": 1, "start_frame": 100, "end_frame_exclusive": 112, "row_start": 0, "row_end_exclusive": 12, "n_samples": 12,
           **{f: None for f in FEATURE_NAMES}, "feature_status": {f: "excluded" for f in FEATURE_NAMES}, "feature_evidence": {}}
    return {"participant_key": "SANI/SANO_A", "session_id": "T1", "timepoint": "T1", "timepoint_kind": "clinical_timepoint",
            "protocol_day": None, "acquisition_date": None, "task": 1, "source_id": source,
            "anatomical_side": side, "limb_condition": condition, "sampling_rate_hz": 60,
            "feature_rows": [row], "flags": [], "status": "partial", "provenance": {"source_sha256": source},
            "channels": {"dtw_hand": {"status": "computed"}},
            "dtw_coordinate_frames": {"1": {"status": "computed", "method_id": DTW_COORDINATE_METHOD,
                                            "global_to_body_row_matrix": np.eye(3).tolist()}}, **metadata}


@pytest.mark.parametrize("field,value", [("participant_key", "CONTROLLO PAZIENTI/SANO_A"), ("session_id", "T2"),
                                        ("timepoint", "T2"), ("protocol_day", 2), ("task", 2), ("acquisition_date", "2030-10-05")])
def test_pairing_never_crosses_identity_or_acquisition_boundary(field, value):
    a, b = pair_result("a", "right"), pair_result("b", "left", **{field: value})
    assert pair_hand_dtw([a, b], {"a": {1: np.zeros((12, 3))}, "b": {1: np.zeros((12, 3))}}) == []
    assert a["feature_rows"][0]["dtw_hand"] is None


def test_ambiguous_pair_is_not_selected_and_patient_pair_is_symmetric():
    a, b = pair_result("a", "right", "paretic"), pair_result("b", "left")
    trajectories = {s: {1: np.arange(36.).reshape(12, 3) * scale} for s, scale in [("a", 1), ("b", 1.1), ("c", 2)]}
    paired = pair_hand_dtw([a, b], trajectories)
    assert len(paired) == 1 and a["feature_rows"][0]["dtw_hand"] == b["feature_rows"][0]["dtw_hand"]
    a, b, c = pair_result("a", "right"), pair_result("b", "left"), pair_result("c", "left")
    assert pair_hand_dtw([a, b, c], trajectories) == []
    assert all(r["flags"][0]["code"] == "dtw_pair_ambiguous" for r in (a, b, c))


def test_simulation_has_explicit_side_and_does_not_replace_healthy_dtw():
    a, b, sim = pair_result("a", "right"), pair_result("b", "left"), pair_result("s", "right", "simulated")
    trajectories = {s: {1: np.arange(36.).reshape(12, 3) * scale} for s, scale in [("a", 1), ("b", 1.1), ("s", 2)]}
    pairs = pair_hand_dtw([a, b, sim], trajectories)
    assert len(pairs) == 2
    assert a["feature_rows"][0]["dtw_hand"] == b["feature_rows"][0]["dtw_hand"]
    assert sim["feature_rows"][0]["dtw_hand"] != b["feature_rows"][0]["dtw_hand"]
    unresolved = pair_result("s", None, "simulated")
    pair_hand_dtw([unresolved, b], trajectories)
    assert unresolved["flags"][0]["code"] == "dtw_pair_metadata_unresolved"


def test_missing_partner_clears_a_previous_distance_when_repairing():
    a, b = pair_result("a", "right"), pair_result("b", "left")
    hand = np.arange(36.).reshape(12, 3)
    trajectories = {"a": {1: hand}, "b": {1: hand * [1, -1, 1]}}
    pair_hand_dtw([a, b], trajectories)
    assert a["feature_rows"][0]["dtw_hand"] == 0
    pair_hand_dtw([a], trajectories)
    assert a["feature_rows"][0]["dtw_hand"] is None
    assert a["feature_rows"][0]["feature_status"]["dtw_hand"] == "excluded"


def test_cli_outputs_all_features_typed_nulls_and_preserves_inputs(full_subject, tmp_path):
    reader = SubjectReader(full_subject)
    paths = list(full_subject.rglob("*")) + list(Path(reader.manifest["source_root"]).rglob("*.xlsx"))
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in paths if p.is_file()}
    output = tmp_path / "features"
    args = ["xsens-features", str(full_subject), "--output", str(output)]
    assert main(args) == 0
    table = pq.read_table(output / "features.parquet")
    assert table.num_rows == 8 and all(table.schema.field(f).type == pa.float64() for f in FEATURE_NAMES)
    report = json.loads((output / "metadata.json").read_text())
    assert report["method"]["environment"]["dtaidistance"]
    assert len(json.loads((output / "dtw_pairs.json").read_text())) == 4
    assert all(hashlib.sha256(p.read_bytes()).hexdigest() == sha for p, sha in before.items())
    assert main([*args, "--anatomical-side", "right"]) == 1
    assert pq.read_table(output / "features.parquet")["dtw_hand"].null_count == 4
    assert main([*args, "--session", "T2"]) == 2
    assert pq.read_table(output / "features.parquet").num_rows == 0
    # A valid configuration incompatible with one recording is a recorded failure,
    # not an invocation failure; the report must still be written.
    assert main([*args, "--sparc-cutoff-hz", "30"]) == 1
    assert json.loads((output / "metadata.json").read_text())["subjects"][0]["recording_status_counts"] == {"failed": 2}
    with pytest.raises(SystemExit) as exc:
        main([*args, "--sparc-cutoff-hz", "0"])
    assert exc.value.code == 2


@pytest.mark.parametrize("marker_count", [4, 6, 7])
def test_invalid_marker_count_skips_entire_recording_and_keeps_other_arm(full_subject, tmp_path, marker_count):
    recording = opened(full_subject)
    raw_root = Path(recording.subject.manifest["source_root"])
    source = raw_root / recording.record["source_path"]
    book = openpyxl.load_workbook(source)
    markers = book["Markers"]
    markers.delete_rows(2, markers.max_row)
    for i in range(marker_count):
        markers.append([1000 + i * 30, f"event{i}"])
    book.save(source)
    write_plan(make_plan(scan(raw_root)), full_subject.parents[1])
    format_subject(full_subject)
    before = opened(full_subject).read_table("Markers").to_pydict()
    output = tmp_path / "strict"
    assert main(["xsens-features", str(full_subject), "--output", str(output), "--no-progress"]) == 1
    report = json.loads((output / "metadata.json").read_text())
    result = next(r for r in report["subjects"][0]["results"] if r["anatomical_side"] == "right")
    assert result["status"] == "skipped" and result["feature_rows"] == 0
    assert result["flags"][0]["code"] == "marker_count_invalid"
    table = pq.read_table(output / "features.parquet")
    assert table.num_rows == 4 and table["anatomical_side"].to_pylist() == ["left"] * 4
    assert table["dtw_hand"].null_count == 4  # Partner is unavailable, never synthesized.
    assert opened(full_subject).read_table("Markers").to_pydict() == before
    assert report["method"]["rom"]["marker_policy_id"] == "2026-10-07-five-markers-strict-v1"


def test_cli_progress_precedes_work_keeps_json_clean_and_preserves_features(full_subject, tmp_path, capsys, monkeypatch):
    import xsens_code.xsens_cli as cli

    extraction = cli.extract_xsens
    observed_starts = []

    def checked_extraction(recording, config):
        captured = capsys.readouterr()
        # The current recording is announced before the potentially slow read/extraction.
        assert "Reading/extracting" in captured.err
        assert captured.out == ""
        observed_starts.append(captured.err)
        return extraction(recording, config)

    monkeypatch.setattr(cli, "extract_xsens", checked_extraction)
    output = tmp_path / "with-progress"
    args = ["xsens-features", str(full_subject), "--output", str(output)]
    assert main(args) == 0
    captured = capsys.readouterr()
    assert len(observed_starts) == 2
    assert "0/2 (0%)" in observed_starts[0] and "1/2 (50%)" in observed_starts[1]
    assert "2/2 (100%)" in captured.err and "DTW pair" in captured.err
    assert "Writing feature tables and metadata" in captured.err and "Finished; exit status 0" in captured.err
    assert "elapsed" in captured.err
    assert json.loads(captured.out)["subjects"][0]["feature_rows"] == 8

    monkeypatch.setattr(cli, "extract_xsens", extraction)
    quiet = tmp_path / "without-progress"
    assert main(["xsens-features", str(full_subject), "--output", str(quiet), "--no-progress"]) == 0
    captured = capsys.readouterr()
    assert captured.err == "" and json.loads(captured.out)["subjects"][0]["feature_rows"] == 8
    assert pq.read_table(output / "features.parquet").equals(pq.read_table(quiet / "features.parquet"))
    assert (output / "features.csv").read_bytes() == (quiet / "features.csv").read_bytes()
    assert (output / "dtw_pairs.json").read_bytes() == (quiet / "dtw_pairs.json").read_bytes()


def test_progress_reports_recording_failures_and_empty_selections(full_subject, tmp_path, capsys):
    args = ["xsens-features", str(full_subject), "--output", str(tmp_path / "features")]
    assert main([*args, "--sparc-cutoff-hz", "30"]) == 1
    captured = capsys.readouterr()
    assert captured.err.count("status=failed") == 2
    assert "2/2 (100%)" in captured.err and "Finished; exit status 1" in captured.err
    assert json.loads(captured.out)["subjects"][0]["recording_status_counts"] == {"failed": 2}

    assert main([*args, "--session", "T2"]) == 2
    captured = capsys.readouterr()
    assert "selected 0 recording(s)" in captured.err and "empty recording selection" in captured.err
    assert json.loads(captured.out)["subjects"][0]["selected_recordings"] == 0
