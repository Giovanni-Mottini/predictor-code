from __future__ import annotations

import ast
import hashlib
import json
from pathlib import Path

import numpy as np
import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from scipy.signal import butter, filtfilt, sosfiltfilt

from predictor_import.adapter import atomic_json, make_plan, write_plan
from predictor_import.formatter import format_subject
from predictor_import.inspection import XSENS_SHEETS
from predictor_import.inventory import scan
from xsens_code.__main__ import main
from shared_code.reader import DataContractError, SubjectReader
from xsens_code.rom import ROM_COLUMNS, ROMConfig, extract_rom, jump_findings, unwrap_angle
from xsens_code.tests.test_xsens_reader import opened, replace_column, rewrite_table


@pytest.fixture
def rom_subject(tmp_path):
    raw = tmp_path / "raw"
    source = raw / "SANI/SANO_FG91/XSens/FG91_t1_task1_dx.xlsx"
    source.parent.mkdir(parents=True)
    book = openpyxl.Workbook()
    book.remove(book.active)
    t = np.arange(120) / 60
    for name in XSENS_SHEETS:
        sheet = book.create_sheet(name)
        if name == "General Information":
            sheet.append(["Frame Rate", 60])
        elif name == "Markers":
            sheet.append(["Frame", "Marker Text"])
            for i, frame in enumerate([1010, 1035, 1060, 1085, 1110]):
                sheet.append([frame, f"event{i}"])
        elif name == "Joint Angles ZXY":
            sheet.append(["Frame", *[f"{side} {suffix}" for side in ("Right", "Left") for suffix in ROM_COLUMNS.values()],
                          "Right Elbow Ulnar Deviation/Radial Deviation", "Left Elbow Ulnar Deviation/Radial Deviation"])
            for i in range(120):
                sheet.append([1000 + i, *[float(30 + (10 + j) * np.sin(2 * np.pi * 0.7 * t[i]) + 2 * np.sin(2 * np.pi * 12 * t[i])) for j in range(2 * len(ROM_COLUMNS))], 0., 0.])
        else:
            sheet.append(["Frame", "Right Hand x", "Left Hand x"])
            for i in range(120):
                sheet.append([1000 + i, i / 10, -i / 10])
    book.save(source)
    output = tmp_path / "standardized"
    write_plan(make_plan(scan(raw)), output)
    subject = output / "SANI/SANO_FG91"
    assert format_subject(subject)["converted"] == 1
    return subject


def legacy_functions():
    root = Path(__file__).resolve().parents[2]
    context = {"np": np, "butter": butter, "filtfilt": filtfilt, "FS_HZ": 60.0}
    for file, name in [("preprocessing.py", "butter_lowpass"), ("features.py", "angular_range")]:
        tree = ast.parse((root / "archive/xsens_code/src" / file).read_text())
        function = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == name)
        exec(compile(ast.Module(body=[function], type_ignores=[]), file, "exec"), context)
    return context["butter_lowpass"], context["angular_range"]


def test_matches_actual_inherited_numerical_functions_and_independent_sos(rom_subject):
    recording = opened(rom_subject)
    table = recording.read_table("Joint Angles ZXY")
    result = extract_rom(recording, ROMConfig(6, (10, 15)))
    legacy_filter, legacy_range = legacy_functions()
    for feature, suffix in ROM_COLUMNS.items():
        raw = np.asarray(table[f"Right {suffix}"].to_pylist())
        expected = legacy_filter(raw, 6)
        independent = sosfiltfilt(butter(4, 6, fs=60, output="sos"), raw, padlen=12)
        for row in result["feature_rows"]:
            start, end = row["row_start"], row["row_end_exclusive"]
            assert row[feature] == legacy_range(expected[start:end])
            assert row[feature] == pytest.approx(np.ptp(independent[start:end]), abs=1e-8)
            assert row["n_samples"] == 25 and row["start_frame"] == 1000 + start
    assert result["computed_feature_values"] == result["qc_pass_feature_values"] == 12
    assert len(result["comparisons"]) == 24
    assert all(r["included_in_summary"] for r in result["comparisons"])
    assert not result["ready_for_features"]


@pytest.mark.parametrize("invalid", [None, float("nan"), float("inf"), -float("inf")])
def test_invalid_sample_outside_windows_excludes_only_affected_full_channel(rom_subject, invalid):
    column = "Right Shoulder Flexion/Extension"
    def change(table):
        values = table[column].to_pylist()
        values[0] = invalid
        return replace_column(table, column, values)
    rewrite_table(rom_subject, "Joint Angles ZXY", change)
    result = extract_rom(opened(rom_subject))
    assert result["status"] == "partial" and result["computed_feature_values"] == 8
    for row in result["feature_rows"]:
        assert row["rom_shoulder_flex_ext"] is None
        assert row["excluded_features"] == ["rom_shoulder_flex_ext"]
        assert all(row[other] is not None for other in list(ROM_COLUMNS)[1:])
    finding = next(f for f in result["channels"]["rom_shoulder_flex_ext"]["findings"] if f["severity"] == "error")
    assert finding["first_frame"] == 1000 and finding["first_source_excel_row"] == 2
    assert all(count == 0 for count in finding["repetition_counts"].values())


def test_bad_other_arm_and_unrelated_table_do_not_exclude_selected_rom(rom_subject):
    rewrite_table(rom_subject, "Joint Angles ZXY", lambda t: replace_column(t, "Left Shoulder Flexion/Extension", [None] * 120))
    recording = opened(rom_subject)
    path = recording.descriptor_path.parent / recording.table_descriptor("Segment Velocity")["file"]
    path.write_bytes(b"not needed by this feature group")
    result = extract_rom(opened(rom_subject))
    assert result["qc_pass_feature_values"] == 12
    assert set(result["provenance"]["checked_tables"]) == {"General Information", "Markers", "Joint Angles ZXY"}


def test_wrapping_is_corrected_before_filtering_and_comparisons_preserve_source(rom_subject):
    column = "Right Shoulder Internal/External Rotation"
    def change(table):
        values = ((np.linspace(170, 200, 120) + 180) % 360 - 180).tolist()
        return replace_column(table, column, values)
    rewrite_table(rom_subject, "Joint Angles ZXY", change)
    result = extract_rom(opened(rom_subject), ROMConfig(6, (10,)))
    channel = result["channels"]["rom_shoulder_int_ext_rot"]
    assert channel["status"] == "computed"
    assert result["computed_feature_values"] == result["qc_pass_feature_values"] == 12
    expected = sosfiltfilt(butter(4, 6, fs=60, output="sos"), np.linspace(170, 200, 120), padlen=12)
    for row in result["feature_rows"]:
        segment = slice(row["row_start"], row["row_end_exclusive"])
        assert row["rom_shoulder_int_ext_rot"] == pytest.approx(np.ptp(expected[segment]), abs=1e-9)
        assert "before_unwrap" in row["feature_evidence"]["rom_shoulder_int_ext_rot"]
        assert row["review_features"] == []
    assert all(r["included_in_summary"] for r in result["comparisons"])
    finding = next(f for f in channel["findings"] if f["code"] == "angle_wrap_corrected")
    assert finding["count"] == 1 and finding["severity"] == "info"
    assert finding["correction_events"][0]["offset_change_deg"] == 360
    assert opened(rom_subject).read_table("Joint Angles ZXY")[column].to_pylist() == ((np.linspace(170, 200, 120) + 180) % 360 - 180).tolist()


def test_raw_jump_screen_retains_original_difference():
    assert jump_findings(np.asarray([0., 180., 0.]), [0, 1, 2], "test", []) == []
    result = jump_findings(np.asarray([179., -179.]), [0, 1], "test", [])
    assert result[0]["count"] == 1 and result[0]["jump_value_repr"] == "-358.0"


@pytest.mark.parametrize("raw,expected", [([178, 179, -179, -178], [178, 179, 181, 182]),
                                        ([-178, -179, 179, 178], [-178, -179, -181, -182]),
                                        ([179, -179, 179, -179], [179, 181, 179, 181])])
def test_unwrap_directions_and_repeated_crossings(raw, expected):
    values = np.asarray(raw, dtype=float)
    corrected, findings = unwrap_angle(values, list(range(len(raw))), "angle", [])
    assert corrected.tolist() == expected
    assert values.tolist() == raw
    assert len(findings[0]["correction_events"]) == findings[0]["count"]


def test_unwrap_preserves_continuous_values_exact_threshold_and_multiple_turns():
    continuous = np.asarray([0., 180., 360., 540., 720.])
    assert unwrap_angle(continuous, list(range(5)), "angle", []) == (continuous, [])
    wrapped = (np.arange(0., 721., 60.) + 180) % 360 - 180
    corrected, findings = unwrap_angle(wrapped, list(range(len(wrapped))), "angle", [])
    assert corrected.tolist() == np.arange(0., 721., 60.).tolist()
    assert findings[0]["count"] == 2


def test_wrap_outside_marker_window_is_corrected_before_full_trace_filter(rom_subject):
    column = "Right Shoulder Internal/External Rotation"
    continuous = np.linspace(178., 238., 120)  # Crossing at row 4; first window starts at row 10.
    rewrite_table(rom_subject, "Joint Angles ZXY", lambda t: replace_column(t, column, ((continuous + 180) % 360 - 180).tolist()))
    result = extract_rom(opened(rom_subject))
    expected = sosfiltfilt(butter(4, 6, fs=60, output="sos"), continuous, padlen=12)
    finding = next(f for f in result["channels"]["rom_shoulder_int_ext_rot"]["findings"] if f["code"] == "angle_wrap_corrected")
    assert all(count == 0 for count in finding["repetition_counts"].values())
    for row in result["feature_rows"]:
        assert row["rom_shoulder_int_ext_rot"] == pytest.approx(np.ptp(expected[row["row_start"]:row["row_end_exclusive"]]), abs=1e-9)


def test_conflicting_unit_and_missing_column_preserve_other_features(rom_subject):
    recording = opened(rom_subject)
    table = recording.table_descriptor("Joint Angles ZXY")
    next(c for c in table["columns"] if c["name"] == "Right Shoulder Flexion/Extension")["unit"] = "rad"
    atomic_json(recording.descriptor_path, recording.descriptor)
    result = extract_rom(opened(rom_subject))
    assert result["computed_feature_values"] == 8
    assert result["channels"]["rom_shoulder_flex_ext"]["findings"][0]["code"] == "angle_unit_conflict"
    recording = opened(rom_subject)
    table = recording.table_descriptor("Joint Angles ZXY")
    column = "Right Shoulder Internal/External Rotation"
    path = recording.descriptor_path.parent / table["file"]
    pq.write_table(pq.read_table(path).drop([column]), path)
    table["columns"] = [c for c in table["columns"] if c["name"] != column]
    table["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    atomic_json(recording.descriptor_path, recording.descriptor)
    result = extract_rom(opened(rom_subject))
    assert result["computed_feature_values"] == 4
    assert result["channels"]["rom_shoulder_int_ext_rot"]["findings"][0]["code"] == "rom_column_missing"


@pytest.mark.parametrize("config", [(0, ()), (float("nan"), ()), (6, (6,)), (6, (10, 10)), (6, (-1,))])
def test_invalid_configuration_rejected(config):
    with pytest.raises(DataContractError):
        ROMConfig(*config)


def test_cutoff_must_be_below_source_nyquist_and_side_explicit(rom_subject):
    with pytest.raises(DataContractError) as err:
        extract_rom(opened(rom_subject), ROMConfig(6, (30,)))
    assert err.value.code == "cutoff_above_nyquist"
    reader = SubjectReader(rom_subject)
    reader.manifest["records"][0]["anatomical_side"] = None
    atomic_json(rom_subject / "manifest.json", reader.manifest)
    with pytest.raises(DataContractError, match="anatomical side"):
        extract_rom(SubjectReader(rom_subject).open_xsens(SubjectReader(rom_subject).manifest["records"][0]))


def test_rom_filter_uses_the_recorded_rate_instead_of_60hz(rom_subject):
    recording = opened(rom_subject)
    raw_root = Path(recording.subject.manifest["source_root"])
    source = raw_root / recording.record["source_path"]
    book = openpyxl.load_workbook(source)
    book["General Information"].cell(1, 2).value = 120
    book.save(source)
    write_plan(make_plan(scan(raw_root)), rom_subject.parents[1])
    assert format_subject(rom_subject)["converted"] == 1
    recording = opened(rom_subject)
    result = extract_rom(recording)
    values = np.asarray(recording.read_table("Joint Angles ZXY")["Right Shoulder Flexion/Extension"].to_pylist())
    expected = sosfiltfilt(butter(4, 6, fs=120, output="sos"), values, padlen=12)
    wrong_rate = sosfiltfilt(butter(4, 6, fs=60, output="sos"), values, padlen=12)
    row = result["feature_rows"][0]
    selection = slice(row["row_start"], row["row_end_exclusive"])
    assert result["sampling_rate_hz"] == 120
    assert row["duration_s"] == 25 / 120
    assert row["rom_shoulder_flex_ext"] == pytest.approx(np.ptp(expected[selection]), abs=1e-8)
    assert not np.isclose(row["rom_shoulder_flex_ext"], np.ptp(wrong_rate[selection]), rtol=0, atol=1e-8)


def test_zero_motion_zero_baseline_and_method_identity(rom_subject):
    rewrite_table(rom_subject, "Joint Angles ZXY", lambda t: replace_column(t, "Right Shoulder Flexion/Extension", [0.] * 120))
    result = extract_rom(opened(rom_subject), ROMConfig(6, (10,)))
    assert all(row["rom_shoulder_flex_ext"] == 0 for row in result["feature_rows"])
    recording = opened(rom_subject)
    pairs = [{"participant_key": recording.subject.participant_key, "source_id": recording.record["source_id"],
              "source_path": recording.record["source_path"], **pair} for pair in result["comparisons"]]
    assert ROMConfig().fingerprint() == ROMConfig().fingerprint()
    assert ROMConfig(6).fingerprint() != ROMConfig(10).fingerprint()
